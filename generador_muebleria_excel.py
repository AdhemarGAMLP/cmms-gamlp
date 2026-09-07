# generador_muebleria_excel.py
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from copy import copy
import os
from datetime import datetime, date

COLUMNAS_OFICIALES = [
    "SECTOR_ACTUAL",
    "DIRECCION_ADMINISTRATIVA",
    "UNIDAD_ORGANIZACIONAL",
    "FECHA_ASIGNACION",
    "TECNICO_INVENTAREADOR",
    "PERSONA_ASIGNADA",
    "CI_ASIGNADO",
    "TIPO_ACTIVO",
    "DESCRIPCION",
    "MODELO",
    "SERIE",
    "DETALLE_TRANSACCION",
    "CODIGO_SISPAM",
    "BERTIN",
    "SAPM",
    "OBSERVACIONES_DE_ASIGNACION",
    "UBICACION",
    "FECHA_INCORPORACION"
]

def obtener_ruta_plantilla_muebleria():
    posibles = [
        os.path.join(os.path.dirname(__file__), "plantillas", "plantilla_muebleria.xlsx"),
        os.path.join("plantillas", "plantilla_muebleria.xlsx"),
        "c:/Users/HP/Desktop/CMMS_GAMLP/plantillas/plantilla_muebleria.xlsx"
    ]
    for p in posibles:
        if os.path.exists(p):
            return p
    return "plantillas/plantilla_muebleria.xlsx"


def exportar_muebleria_excel(lista_muebles, ruta_salida):
    """
    Exporta la lista de activos (muebles y computadoras) escribiendo directamente
    sobre la plantilla oficial plantilla_muebleria.xlsx o creando un libro idéntico.
    """
    ruta_plantilla = obtener_ruta_plantilla_muebleria()
    if os.path.exists(ruta_plantilla):
        wb = openpyxl.load_workbook(ruta_plantilla)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hoja1"
        fill_hdr = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
        font_hdr = Font(name="Calibri", size=10, bold=True, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='A0A0A0'),
            right=Side(style='thin', color='A0A0A0'),
            top=Side(style='thin', color='A0A0A0'),
            bottom=Side(style='thin', color='A0A0A0')
        )
        for col_idx, col_name in enumerate(COLUMNAS_OFICIALES, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = fill_hdr
            cell.font = font_hdr
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    # Estilos para filas de datos
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    font_data = Font(name="Calibri", size=10, color="000000")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Limpiar datos previos si hubieran filas a partir de la fila 2
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    # Escribir filas de datos a partir de la fila 2
    for row_idx, m in enumerate(lista_muebles, start=2):
        valores = [
            str(m.get("sector_actual") or "SALUD"),
            str(m.get("direccion_administrativa") or ""),
            str(m.get("unidad_organizacional") or ""),
            str(m.get("fecha_asignacion") or ""),
            str(m.get("tecnico_inventareador") or ""),
            str(m.get("persona_asignada") or ""),
            str(m.get("ci_asignado") or ""),
            str(m.get("tipo_activo") or ""),
            str(m.get("descripcion") or ""),
            str(m.get("modelo") or ""),
            str(m.get("serie") or ""),
            str(m.get("detalle_transaccion") or "ASIGNACION"),
            str(m.get("codigo_sispam") or ""),
            str(m.get("bertin") or ""),
            str(m.get("sapm") or ""),
            str(m.get("observaciones_de_asignacion") or ""),
            str(m.get("ubicacion") or ""),
            str(m.get("fecha_incorporacion") or "")
        ]

        for col_idx, val in enumerate(valores, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = font_data
            c.border = thin_border
            # Alineación centrada para códigos, fechas, sectores; izquierda para descripciones
            if col_idx in (1, 4, 7, 10, 11, 12, 13, 14, 15, 18):
                c.alignment = align_center
            else:
                c.alignment = align_left

        ws.row_dimensions[row_idx].height = 20

    # Autoajustar anchos de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws.row_dimensions[1].height = 26
    wb.save(ruta_salida)
    return True, f"Reporte exportado exitosamente a:\n{ruta_salida}"


def importar_muebleria_excel(ruta_archivo):
    """
    Lee un archivo Excel y retorna una lista de diccionarios normalizados
    con las 18 columnas oficiales de Mueblería y Computadoras.
    """
    if not os.path.exists(ruta_archivo):
        return [], 0, "El archivo especificado no existe."

    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        ws = wb.active
    except Exception as e:
        return [], 0, f"Error al abrir archivo Excel: {e}"

    # 1. Detectar fila de encabezados
    header_row_idx = 1
    col_map = {}

    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(cell.value or "").strip().upper().replace(" ", "_") for cell in ws[r]]
        if "SECTOR_ACTUAL" in row_vals or "TIPO_ACTIVO" in row_vals or "CODIGO_SISPAM" in row_vals:
            header_row_idx = r
            for c_idx, val in enumerate(row_vals, start=1):
                if val:
                    col_map[val] = c_idx
            break

    # Si no encontró por nombre exacto, mapear por posición estándar 1..18
    if not col_map:
        for idx, col_name in enumerate(COLUMNAS_OFICIALES, start=1):
            col_map[col_name] = idx

    lista_resultado = []
    total_filas = 0

    for r in range(header_row_idx + 1, ws.max_row + 1):
        def _get_val(col_key, default=""):
            c_pos = col_map.get(col_key)
            if c_pos and c_pos <= ws.max_column:
                v = ws.cell(row=r, column=c_pos).value
                if v is not None:
                    if isinstance(v, (datetime, date)):
                        return v.strftime("%Y-%m-%d")
                    return str(v).strip()
            return default

        sector = _get_val("SECTOR_ACTUAL", "SALUD")
        direccion = _get_val("DIRECCION_ADMINISTRATIVA", "")
        unidad = _get_val("UNIDAD_ORGANIZACIONAL", "")
        tipo_activo = _get_val("TIPO_ACTIVO", "")
        descripcion = _get_val("DESCRIPCION", "")
        modelo = _get_val("MODELO", "")
        serie = _get_val("SERIE", "")
        sispam = _get_val("CODIGO_SISPAM", "")
        bertin = _get_val("BERTIN", "")
        sapm = _get_val("SAPM", "")

        # Verificar si la fila está completamente vacía
        if not any([sector, direccion, unidad, tipo_activo, descripcion, modelo, serie, sispam, bertin, sapm]):
            continue

        item = {
            "sector_actual": sector or "SALUD",
            "direccion_administrativa": direccion,
            "unidad_organizacional": unidad,
            "fecha_asignacion": _get_val("FECHA_ASIGNACION", ""),
            "tecnico_inventareador": _get_val("TECNICO_INVENTAREADOR", ""),
            "persona_asignada": _get_val("PERSONA_ASIGNADA", ""),
            "ci_asignado": _get_val("CI_ASIGNADO", ""),
            "tipo_activo": tipo_activo or "COMPUTADORA",
            "descripcion": descripcion,
            "modelo": modelo,
            "serie": serie,
            "detalle_transaccion": _get_val("DETALLE_TRANSACCION", "ASIGNACION"),
            "codigo_sispam": sispam,
            "bertin": bertin,
            "sapm": sapm,
            "observaciones_de_asignacion": _get_val("OBSERVACIONES_DE_ASIGNACION", ""),
            "ubicacion": _get_val("UBICACION", ""),
            "fecha_incorporacion": _get_val("FECHA_INCORPORACION", ""),
            "estado": "Activo"
        }

        lista_resultado.append(item)
        total_filas += 1

    return lista_resultado, total_filas, None
