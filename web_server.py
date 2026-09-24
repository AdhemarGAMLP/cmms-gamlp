import threading
import os
import io
import json
import base64
import psycopg2.extras
import openpyxl
from openpyxl.styles import Font, Alignment
try:
    import pythoncom
    import win32com.client
except ImportError:
    pythoncom = None
    win32com = None
from flask import Flask, render_template_string, request, redirect, url_for, send_from_directory, send_file, session, jsonify
from functools import wraps
from database import (
    obtener_conexion,
    comprimir_imagen_base64,
    obtener_jerarquia_sedes_db,
    generar_codigo_red,
    generar_sigla_centro,
    generar_siguiente_codigo_af,
    guardar_mueble_db,
    guardar_equipo_db,
    obtener_areas_db,
    obtener_catalogo_equipos_db,
    obtener_equipos_db,
    obtener_muebles_db,
    guardar_area_db,
    guardar_catalogo_db,
    obtener_repuestos_db,
    guardar_repuesto_db,
    obtener_usuarios_db,
    guardar_usuario_permisos_db,
    obtener_intervenciones_db,
    guardar_intervencion_db,
    obtener_estadisticas_censo_db,
    eliminar_registro_db,
    obtener_papelera_db,
    recuperar_registro_papelera_db,
    purgar_registro_papelera_db
)
from auth import login as auth_login
from vistas_web_movil import HTML_MOVIL_LOGIN, HTML_MOVIL_REGISTRO
from vistas_web_historico import HTML_HISTORICO_WEB
from datetime import date, datetime
from excel_utils import (
    obtener_ruta_plantilla,
    escribir_en_celda_segura,
    marcar_x,
    escribir_texto_largo,
    exportar_excel_a_pdf,
    sanitizar_nombre_archivo
)
from config import CARPETAS, CONFIG

app_web = Flask(__name__)
app_web.secret_key = os.environ.get("FLASK_SECRET_KEY", "gamlp_sgem_secret_key_2026_super_secure")
app = app_web  # Alias para servidores WSGI de producción (Gunicorn / Render / Vercel)
app_gui = None  # Referencia global de la GUI de Tkinter para sincronización

def login_requerido(f):
    @wraps(f)
    def decorada(*args, **kwargs):
        if 'usuario' not in session:
            return redirect(url_for('movil_login', next=request.path))
        return f(*args, **kwargs)
    return decorada


# Pantalla de éxito responsiva con enlace de descarga Excel
HTML_EXITO = """
<!DOCTYPE html><html lang="es"><head><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Éxito</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<style>
    * { box-sizing: border-box; }
    body { font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #F8FAFC; margin: 0; padding: 20px; display: flex; align-items: center; justify-content: center; min-height: 100vh; text-align: center; color: #0F172A; -webkit-font-smoothing: antialiased; }
    .tarjeta { background: white; padding: 32px 24px; border-radius: 14px; border: 1px solid #E2E8F0; box-shadow: 0 1px 4px rgba(0,0,0,0.04); max-width: 450px; width: 100%; }
    .icono { font-size: 46px; color: #10B981; margin-bottom: 15px; }
    .btn { display: block; text-decoration: none; padding: 12px 20px; border-radius: 8px; font-weight: 700; margin-top: 10px; font-size: 14px; text-align: center; transition: all 0.2s; }
    .btn-download-xlsx { background: #0F172A; color: white; }
    .btn-download-xlsx:hover { background: #1E293B; }
    .btn-back { background: #FFFFFF; color: #0F172A; border: 1.5px solid #E2E8F0; }
    .btn-back:hover { border-color: #0F172A; }
</style>
<script>
    window.addEventListener('DOMContentLoaded', () => {
        var dlBtn = document.getElementById('btn-auto-dl');
        if (dlBtn) {
            setTimeout(() => {
                dlBtn.click();
            }, 600);
        }
    });
</script>
</head><body>
    <div class="tarjeta">
        <div class="icono">✓</div>
        <h2 style="margin-top:0; font-size: 20px; font-weight: 800;">¡Mantenimiento Guardado!</h2>
        <p style="color: #64748B; font-size: 13.5px; margin-bottom: 24px; line-height: 1.5;">La intervención se registró correctamente y la Hoja de Trabajo se generó en la planilla oficial de Excel.</p>
        
        {% if xlsx_file %}
        <a id="btn-auto-dl" href="/descargar/{{ xlsx_file }}" class="btn btn-download-xlsx">📥 Descargar Hoja de Trabajo (.xlsx)</a>
        {% endif %}
        <a href="/equipo/{{ id_equipo }}" class="btn btn-back">Volver al Equipo</a>
    </div>
</body></html>
"""

def generar_excel_ht_web(eq_data, form_data, realizado_por, sello_firma_path=None):
    ruta_plantilla_ht = obtener_ruta_plantilla("plantilla_trabajo.xlsx")
    if not os.path.exists(ruta_plantilla_ht):
        print(f"[ERROR] No se encontró la plantilla en: {ruta_plantilla_ht}")
        return None, None
        
    try:
        wb = openpyxl.load_workbook(ruta_plantilla_ht)
        ws = wb.active
        
        # Escribir campos de cabecera y equipo
        escribir_en_celda_segura(ws, 'P5', eq_data.get('red_salud_nombre', ''))
        escribir_en_celda_segura(ws, 'P6', eq_data.get('centro_salud_nombre', ''))
        escribir_en_celda_segura(ws, 'F11', eq_data.get('area', ''))
        escribir_en_celda_segura(ws, 'AA11', eq_data.get('servicio', ''))
        escribir_en_celda_segura(ws, 'J15', eq_data.get('nombre', ''))
        escribir_en_celda_segura(ws, 'AE15', str(eq_data.get('id', '')))
        escribir_en_celda_segura(ws, 'E17', eq_data.get('procedencia', ''))
        escribir_en_celda_segura(ws, 'AB17', str(eq_data.get('anio_fab', '')))
        escribir_en_celda_segura(ws, 'E19', eq_data.get('marca', ''))
        escribir_en_celda_segura(ws, 'AB19', eq_data.get('fabricante', ''))
        escribir_en_celda_segura(ws, 'F21', eq_data.get('modelo', ''))
        escribir_en_celda_segura(ws, 'AB21', eq_data.get('numero_serie', ''))
        
        # Fechas
        f_rec_raw = form_data.get('fecha_recepcion', date.today().strftime('%Y-%m-%d'))
        f_ent_raw = form_data.get('fecha_entrega', date.today().strftime('%Y-%m-%d'))
        h_ejec = form_data.get('hora_ejecucion', datetime.now().strftime('%H:%M'))
        try:
            f_rec_dt = datetime.strptime(f_rec_raw, '%Y-%m-%d').date()
            f_rec_str = f_rec_dt.strftime('%d / %m / %Y')
        except:
            f_rec_str = datetime.now().strftime('%d / %m / %Y')
            
        try:
            f_ent_dt = datetime.strptime(f_ent_raw, '%Y-%m-%d').date()
            f_ent_str = f_ent_dt.strftime('%d / %m / %Y')
        except:
            f_ent_str = datetime.now().strftime('%d / %m / %Y')
            
        escribir_en_celda_segura(ws, 'M23', f_rec_str)
        escribir_en_celda_segura(ws, 'I62', f"{f_ent_str}  {h_ejec}")
        
        # Nombre del técnico firmante responsable
        escribir_en_celda_segura(ws, 'J64', realizado_por)
        
        # Condición
        cond = form_data.get('condicion')
        if cond == "Óptimo": marcar_x(ws, 'P26')
        elif cond == "Aceptable": marcar_x(ws, 'W26')
        elif cond == "Crítica": marcar_x(ws, 'AC26')
        elif cond == "Inoperante": marcar_x(ws, 'AJ26')
        elif cond == "F/Servicio": marcar_x(ws, 'AP26')

        # Estado Físico
        est = form_data.get('estado_equipo')
        if est == "Óptimo": marcar_x(ws, 'O29')
        elif est == "Bueno": marcar_x(ws, 'U29')
        elif est == "Regular": marcar_x(ws, 'AB29')
        elif est == "Malo": marcar_x(ws, 'AH29')
        elif est == "Obsoleto": marcar_x(ws, 'AO29')

        # Tipo Mantenimiento
        tipo = form_data.get('tipo')
        if tipo == "Preventivo": marcar_x(ws, 'Q43')
        else: marcar_x(ws, 'AL43')

        # Textos largos
        escribir_texto_largo(ws, 'B33', form_data.get('deficiencia', ''))
        escribir_texto_largo(ws, 'B47', form_data.get('trabajo', ''))
        escribir_texto_largo(ws, 'B53', form_data.get('observaciones', ''))
        
        # Inyectar sello/firma (Imagen) si existe
        if sello_firma_path and os.path.exists(sello_firma_path):
            try:
                from openpyxl.drawing.image import Image as ExcelImage
                img = ExcelImage(sello_firma_path)
                img.width = 145
                img.height = 65
                ws.add_image(img, 'AD60')
            except Exception as ex:
                print(f"[ERROR] No se pudo insertar la firma en Excel: {ex}")
        
        # Nombres de salida
        try:
            f_rec_dt = datetime.strptime(f_rec_raw, '%Y-%m-%d').date()
            fecha_compacta = f_rec_dt.strftime('%Y%m%d')
        except:
            fecha_compacta = datetime.now().strftime('%Y%m%d')
            
        timestamp_seguro = datetime.now().strftime('%H%M%S')
        id_eq_limpio = sanitizar_nombre_archivo(eq_data.get('id', 'EQ'))
        nombre_base = f"HT_{id_eq_limpio}_{fecha_compacta}_{timestamp_seguro}"
        filename_xlsx = f"{nombre_base}.xlsx"
        filename_pdf = f"{nombre_base}.pdf"
        
        area_name = eq_data.get("area", "General") or "General"
        area_folder = sanitizar_nombre_archivo(area_name)
        dir_mantenimiento = os.path.join(CARPETAS["areas"], area_folder, "mantenimientos")
        os.makedirs(dir_mantenimiento, exist_ok=True)
        
        ruta_xlsx = os.path.join(dir_mantenimiento, filename_xlsx)
        ruta_pdf = os.path.join(dir_mantenimiento, filename_pdf)
        
        wb.save(ruta_xlsx)
        
        # Intentar renderizar a PDF
        exportar_excel_a_pdf(ruta_xlsx, ruta_pdf, rango_impresion="$A$1:$AR$67")
        
        return filename_xlsx, filename_pdf
    except Exception as e:
        print(f"[ERROR] Fallo al generar archivos de Hoja de Trabajo: {e}")
        return None, None

HTML_INVENTARIO = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Inventario Biomédico | SGEM GAMLP</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #0F172A;
            --primary-dark: #020617;
            --primary-hover: #1E293B;
            --bg: #F8FAFC;
            --card: #FFFFFF;
            --text: #0F172A;
            --muted: #64748B;
            --border: #E2E8F0;
            --border-hover: #CBD5E1;
            --success: #10B981;
            --warning: #F59E0B;
            --danger: #EF4444;
            --purple: #8B5CF6;
        }
        * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; }
        body { background: var(--bg); color: var(--text); padding-bottom: 50px; -webkit-font-smoothing: antialiased; }
        
        .header {
            background: #0F172A;
            color: white;
            padding: 24px 20px 20px;
            text-align: center;
            border-bottom: 1px solid #1E293B;
            box-shadow: 0 1px 3px rgba(0,0,0,0.1);
        }
        .header h1 { font-size: 22px; font-weight: 800; margin-bottom: 4px; letter-spacing: -0.5px; }
        .header p { font-size: 13px; color: #94A3B8; }
        .badge-gamlp { display: inline-block; background: rgba(255,255,255,0.08); border: 1px solid rgba(255,255,255,0.14); color: #CBD5E1; padding: 3px 12px; border-radius: 20px; font-size: 11px; margin-bottom: 8px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.5px; }

        .container { max-width: 960px; margin: -15px auto 0; padding: 0 16px; }
        
        .filter-card {
            background: var(--card);
            border-radius: 14px;
            padding: 18px;
            box-shadow: 0 1px 4px rgba(0, 0, 0, 0.04);
            border: 1px solid var(--border);
            margin-bottom: 20px;
        }
        .filter-grid {
            display: grid;
            grid-template-columns: 1fr;
            gap: 12px;
            margin-bottom: 12px;
        }
        @media (min-width: 640px) {
            .filter-grid { grid-template-columns: 1fr 1fr; }
        }
        @media (min-width: 860px) {
            .filter-grid { grid-template-columns: 1.2fr 1.2fr 1fr; }
        }

        .filter-group { display: flex; flex-direction: column; gap: 4px; }
        .filter-label { font-size: 11px; font-weight: 700; color: var(--muted); text-transform: uppercase; letter-spacing: 0.5px; }
        
        .select-input, .search-input {
            width: 100%;
            padding: 11px 14px;
            border-radius: 8px;
            border: 1px solid var(--border);
            font-size: 14px;
            outline: none;
            background: #F8FAFC;
            color: var(--text);
            font-weight: 500;
            transition: all 0.2s;
        }
        .select-input:focus, .search-input:focus { border-color: var(--border-hover); background: #FFFFFF; box-shadow: 0 0 0 3px rgba(15,23,42,0.08); }
        
        .search-input { background: #FFFFFF; }

        .stats-grid {
            display: grid;
            grid-template-columns: repeat(2, 1fr);
            gap: 10px;
            margin-bottom: 20px;
        }
        @media (min-width: 600px) {
            .stats-grid { grid-template-columns: repeat(4, 1fr); }
        }
        .stat-card {
            background: var(--card);
            padding: 14px 16px;
            border-radius: 12px;
            border: 1px solid var(--border);
            box-shadow: 0 1px 3px rgba(0,0,0,0.02);
            text-align: center;
        }
        .stat-num { font-size: 26px; font-weight: 800; color: var(--text); letter-spacing: -0.5px; }
        .stat-lbl { font-size: 11px; color: var(--muted); font-weight: 600; margin-top: 2px; }

        .area-section {
            background: var(--card);
            border-radius: 14px;
            padding: 18px;
            border: 1px solid var(--border);
            margin-bottom: 16px;
            box-shadow: 0 1px 3px rgba(0,0,0,0.02);
        }
        .area-header {
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding-bottom: 12px;
            border-bottom: 1px solid var(--border);
            margin-bottom: 14px;
        }
        .area-title {
            font-size: 15px;
            font-weight: 700;
            color: var(--text);
            display: flex;
            align-items: center;
            gap: 8px;
        }
        .area-badge-count {
            background: #F1F5F9;
            color: #334155;
            font-size: 11px;
            font-weight: 700;
            padding: 3px 10px;
            border-radius: 20px;
            border: 1px solid #E2E8F0;
        }

        .equipment-list { display: flex; flex-direction: column; gap: 10px; }
        .equipment-card {
            background: #FFFFFF;
            border-radius: 10px;
            padding: 14px 16px;
            border: 1px solid var(--border);
            box-shadow: 0 1px 2px rgba(0,0,0,0.02);
            transition: all 0.15s ease-in-out;
            text-decoration: none;
            color: inherit;
            display: block;
        }
        .equipment-card:hover { border-color: #CBD5E1; box-shadow: 0 3px 10px rgba(0,0,0,0.04); }
        .equipment-card:active { transform: scale(0.99); }
        
        .eq-header { display: flex; justify-content: space-between; align-items: flex-start; gap: 8px; margin-bottom: 6px; }
        .eq-title { font-size: 14.5px; font-weight: 700; color: var(--text); }
        .eq-code { background: #F1F5F9; color: #334155; border: 1px solid #E2E8F0; font-size: 11px; font-weight: 700; padding: 2px 8px; border-radius: 6px; white-space: nowrap; }
        
        .eq-detail { font-size: 13px; color: var(--muted); margin-bottom: 4px; display: flex; align-items: center; gap: 6px; }
        .eq-badges { display: flex; flex-wrap: wrap; gap: 5px; margin-top: 8px; }
        .badge { font-size: 11px; font-weight: 600; padding: 3px 7px; border-radius: 6px; }
        .badge-red { background: #F8FAFC; color: #0284C7; border: 1px solid #E2E8F0; }
        .badge-centro { background: #F8FAFC; color: #334155; border: 1px solid #E2E8F0; }
        .badge-area { background: #F8FAFC; color: #475569; border: 1px solid #E2E8F0; }
        .badge-garantia { background: #ECFDF5; color: #065F46; border: 1px solid #A7F3D0; }
        .badge-mtto { background: #FEF3C7; color: #92400E; border: 1px solid #FDE68A; }
        .badge-danger { background: #FEE2E2; color: #991B1B; border: 1px solid #FECACA; }

        .btn-view {
            display: inline-block;
            margin-top: 8px;
            font-size: 12px;
            font-weight: 700;
            color: #0F172A;
        }
        .no-results {
            text-align: center;
            padding: 40px 20px;
            color: var(--muted);
            font-size: 15px;
            font-weight: 600;
            display: none;
        }
        .asset-filters-bar {
            display: flex;
            gap: 8px;
            margin-bottom: 20px;
            flex-wrap: wrap;
            background: #FFFFFF;
            padding: 6px;
            border-radius: 12px;
            border: 1px solid var(--border);
            box-shadow: 0 1px 3px rgba(0,0,0,0.03);
        }
        .asset-pill-btn {
            background: transparent;
            border: 1.5px solid transparent;
            color: #64748B;
            padding: 8px 18px;
            border-radius: 8px;
            font-size: 13.5px;
            font-weight: 700;
            cursor: pointer;
            transition: all 0.2s ease-in-out;
            display: inline-flex;
            align-items: center;
            gap: 8px;
        }
        .asset-pill-btn:hover {
            background: #F8FAFC;
            color: #0F172A;
            border-color: #E2E8F0;
        }
        .asset-pill-btn.active {
            background: #007AFF;
            border-color: #007AFF;
            color: #FFFFFF;
            box-shadow: 0 2px 8px rgba(0, 122, 255, 0.3);
        }
        .nav-tabs {
            display: flex;
            justify-content: center;
            gap: 8px;
            margin-top: 15px;
        }
        .nav-tab {
            color: #94A3B8;
            text-decoration: none;
            padding: 6px 14px;
            border-radius: 20px;
            font-size: 12.5px;
            font-weight: 600;
            background: rgba(255, 255, 255, 0.06);
            border: 1px solid rgba(255, 255, 255, 0.1);
            transition: all 0.2s;
        }
        .nav-tab:hover {
            color: #FFFFFF;
            border-color: rgba(255, 255, 255, 0.25);
        }
        .nav-tab.active {
            background: #FFFFFF;
            color: #0F172A;
            font-weight: 700;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.2);
        }
    </style>
</head>
<body>
    <div class="header">
        <span class="badge-gamlp">GAMLP • SGEM v1.1</span>
        <h1>Sistema de Gestión de Equipamiento Médico</h1>
        <p>Inventario Descentralizado por Redes y Centros de Salud</p>
        <div class="nav-tabs">
            <a href="/inventario" class="nav-tab active">📦 Inventario</a>
            <a href="/analisis" class="nav-tab">📊 Análisis y Censo</a>
            <a href="/movil" class="nav-tab">📱 Registro Móvil</a>
            <a href="/historico" class="nav-tab">🏛️ Histórico</a>
            <a href="/enlaces" class="nav-tab">🔗 Enlaces</a>
        </div>
    </div>

    <div class="container">
        <!-- Tarjeta de Filtros en Cascada -->
        <div class="filter-card">
            <div class="filter-grid">
                <div class="filter-group">
                    <label class="filter-label">🌐 Red de Salud</label>
                    <select id="filtro-red" class="select-input" onchange="alCambiarRed()">
                        <option value="">Todas las Redes (GAMLP)</option>
                        {% for r in redes %}
                        <option value="{{ r['nombre'] }}" data-id="{{ r['id'] }}">{{ r['nombre'] }}</option>
                        {% endfor %}
                    </select>
                </div>
                <div class="filter-group">
                    <label class="filter-label">🏥 Centro de Salud</label>
                    <select id="filtro-centro" class="select-input" onchange="alCambiarCentro()">
                        <option value="">Todos los Centros</option>
                        {% for c in centros %}
                        <option value="{{ c['nombre'] }}" data-red-id="{{ c['red_salud_id'] }}">{{ c['nombre'] }}</option>
                        {% endfor %}
                    </select>
                </div>
                <div class="filter-group">
                    <label class="filter-label">📂 Área / Servicio</label>
                    <select id="filtro-area" class="select-input" onchange="filtrar()">
                        <option value="">Todas las Áreas</option>
                        {% for a in areas %}
                        <option value="{{ a }}">{{ a }}</option>
                        {% endfor %}
                    </select>
                </div>
            </div>
            
            <div class="filter-group">
                <input type="text" id="busqueda" class="search-input" placeholder="🔍 Buscar por nombre, marca, serie, código AF o servicio..." onkeyup="filtrar()">
            </div>
        </div>

        <!-- Pestañas de Selección de Vista (Ambos, Equipos Médicos, Activos Fijos) -->
        <div class="asset-filters-bar">
            <button type="button" class="asset-pill-btn active" data-tipo="TODO" onclick="filtrarTipoActivo('TODO')">
                🌐 Ver Ambos / Todo (<span id="cnt-pills-todo">{{ cnt_todo }}</span>)
            </button>
            <button type="button" class="asset-pill-btn" data-tipo="EQUIPO" onclick="filtrarTipoActivo('EQUIPO')">
                🩺 Solo Equipos Médicos (<span id="cnt-pills-equipos">{{ cnt_equipos }}</span>)
            </button>
            <button type="button" class="asset-pill-btn" data-tipo="MUEBLE" onclick="filtrarTipoActivo('MUEBLE')">
                🛋️ Solo Activos Fijos (<span id="cnt-pills-muebles">{{ cnt_muebles }}</span>)
            </button>
        </div>

        <!-- Tarjetas de Estadísticas -->
        <div class="stats-grid">
            <div class="stat-card">
                <div class="stat-num" id="stat-total">{{ total }}</div>
                <div class="stat-lbl" id="lbl-stat-total">Total Activos</div>
            </div>
            <div class="stat-card">
                <div class="stat-num" style="color: var(--success);" id="stat-op">{{ operativos }}</div>
                <div class="stat-lbl" id="lbl-stat-op">Operativos / Buen Estado</div>
            </div>
            <div class="stat-card">
                <div class="stat-num" style="color: var(--warning);" id="stat-gar">{{ garantia }}</div>
                <div class="stat-lbl" id="lbl-stat-gar">En Garantía / Regular</div>
            </div>
            <div class="stat-card">
                <div class="stat-num" style="color: var(--danger);" id="stat-baj">{{ bajas }}</div>
                <div class="stat-lbl" id="lbl-stat-baj">Bajas / Fuera de Servicio</div>
            </div>
        </div>

        <!-- Contenedor de Áreas y Equipos -->
        <div id="contenedor-areas">
            {% for area_nom in areas %}
            <div class="area-section" data-area-name="{{ area_nom }}">
                <div class="area-header">
                    <div class="area-title">
                        <span>🚪 Área: {{ area_nom }}</span>
                    </div>
                    <span class="area-badge-count count-label">0 activos</span>
                </div>
                <div class="equipment-list">
                    {% for ac in activos %}
                    {% if (ac['area'] or 'General') == area_nom %}
                    {% if ac['_tipo'] == 'EQUIPO' %}
                    <a href="/equipo/{{ ac['id'] }}" 
                       class="equipment-card" 
                       data-tipo="EQUIPO"
                       data-red="{{ ac['red_salud_nombre'] or '' }}"
                       data-centro="{{ ac['centro_salud_nombre'] or '' }}"
                       data-area="{{ ac['area'] or '' }}"
                       data-estado="{{ ac['estado'] or '' }}"
                       data-garantia="{{ ac['garantia'] or '' }}"
                       data-texto="{{ ac['nombre'] }} {{ ac['marca'] }} {{ ac['modelo'] }} {{ ac['id'] }} {{ ac['numero_serie'] }} {{ ac['servicio'] }} {{ ac['area'] }} {{ ac['red_salud_nombre'] }} {{ ac['centro_salud_nombre'] }}">
                        <div class="eq-header">
                            <div style="display: flex; align-items: center; gap: 10px;">
                                {% if ac['foto'] %}
                                <img src="{{ ac['foto'] }}" alt="" style="width: 44px; height: 44px; border-radius: 8px; object-fit: cover; border: 1px solid #E2E8F0; flex-shrink: 0;">
                                {% else %}
                                <span style="font-size: 24px;">🩺</span>
                                {% endif %}
                                <div class="eq-title">{{ ac['nombre'] }}</div>
                            </div>
                            <span class="eq-code">{{ ac['id'] }}</span>
                        </div>
                        <div class="eq-detail">🏷️ <strong>{{ ac['marca'] }}</strong> - {{ ac['modelo'] }} {% if ac['numero_serie'] and ac['numero_serie'] != 'S/N' %}| S/N: {{ ac['numero_serie'] }}{% endif %}</div>
                        <div class="eq-detail">📍 Servicio: <strong>{{ ac['servicio'] or ac['area'] }}</strong></div>
                        <div class="eq-badges">
                            <span class="badge" style="background: #EFF6FF; color: #1D4ED8; font-weight: 700;">🩺 Equipo Médico</span>
                            {% if ac['red_salud_nombre'] %}
                                <span class="badge badge-red">🌐 {{ ac['red_salud_nombre'] }}</span>
                            {% endif %}
                            {% if ac['centro_salud_nombre'] %}
                                <span class="badge badge-centro">🏥 {{ ac['centro_salud_nombre'] }}</span>
                            {% endif %}
                            {% if ac['garantia'] == 'Con Garantía' %}
                                <span class="badge badge-garantia">🛡️ Con Garantía</span>
                            {% endif %}
                            {% if ac['f_prox'] %}
                                <span class="badge badge-mtto">📅 Próx. Mtto: {{ ac['f_prox'] }}</span>
                            {% endif %}
                            {% if ac['estado'] == 'Baja' %}
                                <span class="badge badge-danger">Dado de Baja</span>
                            {% endif %}
                        </div>
                        <div class="btn-view">Ver Ficha Técnica Completa →</div>
                    </a>
                    {% else %}
                    <!-- TARJETA DE MUEBLE / COMPUTADORA -->
                    <a href="/mueble/{{ ac['id_db'] }}" 
                       class="equipment-card" 
                       data-tipo="MUEBLE"
                       data-red="{{ ac['red_salud_nombre'] or '' }}"
                       data-centro="{{ ac['centro_salud_nombre'] or '' }}"
                       data-area="{{ ac['area'] or '' }}"
                       data-estado="{{ ac['estado'] or '' }}"
                       data-garantia=""
                       data-texto="{{ ac['nombre'] }} {{ ac['tipo_activo'] }} {{ ac['marca'] }} {{ ac['modelo'] }} {{ ac['codigo_af'] }} {{ ac['numero_serie'] }} {{ ac['servicio'] }} {{ ac['area'] }} {{ ac['red_salud_nombre'] }} {{ ac['centro_salud_nombre'] }}">
                        <div class="eq-header">
                            <div style="display: flex; align-items: center; gap: 10px;">
                                <span style="font-size: 24px;">{{ ac['_icono'] }}</span>
                                <div class="eq-title">{{ ac['nombre'] }}</div>
                            </div>
                            <span class="eq-code" style="background: #FEF3C7; color: #92400E;">{{ ac['codigo_af'] }}</span>
                        </div>
                        <div class="eq-detail">🏷️ <strong>{{ ac['marca'] }}</strong> - {{ ac['modelo'] }} {% if ac['numero_serie'] and ac['numero_serie'] != 'S/C' and ac['numero_serie'] != 'S/N' %}| S/N: {{ ac['numero_serie'] }}{% endif %}</div>
                        <div class="eq-detail">📍 Ubicación: <strong>{{ ac['servicio'] or ac['area'] }}</strong></div>
                        <div class="eq-badges">
                            <span class="badge" style="background: #FEF3C7; color: #92400E; font-weight: 700;">{{ ac['_icono'] }} {{ ac['tipo_activo'] or 'Mueble / TI' }}</span>
                            {% if ac['red_salud_nombre'] %}
                                <span class="badge badge-red">🌐 {{ ac['red_salud_nombre'] }}</span>
                            {% endif %}
                            {% if ac['centro_salud_nombre'] %}
                                <span class="badge badge-centro">🏥 {{ ac['centro_salud_nombre'] }}</span>
                            {% endif %}
                            {% if ac['estado'] == 'Baja' %}
                                <span class="badge badge-danger">Dado de Baja</span>
                            {% elif ac['estado'] == 'Regular' %}
                                <span class="badge badge-mtto">Estado Regular</span>
                            {% else %}
                                <span class="badge badge-garantia">Operativo / Bueno</span>
                            {% endif %}
                            {% if ac['responsable'] %}
                                <span class="badge" style="background: #F1F5F9; color: #475569;">👤 {{ ac['responsable'] }}</span>
                            {% endif %}
                        </div>
                        <div class="btn-view" style="color: #D97706;">Ver Ficha de Mueble / TI →</div>
                    </a>
                    {% endif %}
                    {% endif %}
                    {% endfor %}
                </div>
            </div>
            {% endfor %}
        </div>

        <div id="no-results" class="no-results">
            🔍 No se encontraron activos con los filtros seleccionados.
        </div>
    </div>

    <script>
        // Mapeo dinámico de áreas por centro de salud
        const mapaAreasPorCentro = {{ mapa_areas_centro | safe }};
        const todasLasAreasOriginales = {{ areas | tojson }};

        // Lista original de centros para filtrado en cascada
        const todosLosCentros = Array.from(document.querySelectorAll('#filtro-centro option')).map(opt => ({
            value: opt.value,
            text: opt.text,
            redId: opt.getAttribute('data-red-id')
        }));

        let TIPO_ACTIVO_FILTRO = 'TODO';

        function filtrarTipoActivo(tipo) {
            TIPO_ACTIVO_FILTRO = tipo;
            document.querySelectorAll('.asset-pill-btn').forEach(btn => {
                if (btn.getAttribute('data-tipo') === tipo) {
                    btn.classList.add('active');
                } else {
                    btn.classList.remove('active');
                }
            });
            filtrar();
        }

        function alCambiarRed() {
            const selectRed = document.getElementById('filtro-red');
            const selectCentro = document.getElementById('filtro-centro');
            const redOpt = selectRed.options[selectRed.selectedIndex];
            const redId = redOpt ? redOpt.getAttribute('data-id') : '';

            // Filtrar centros de salud según la red
            selectCentro.innerHTML = '<option value="">Todos los Centros</option>';
            todosLosCentros.forEach(c => {
                if (c.value === '') return;
                if (!redId || c.redId === redId) {
                    const opt = document.createElement('option');
                    opt.value = c.value;
                    opt.textContent = c.text;
                    opt.setAttribute('data-red-id', c.redId);
                    selectCentro.appendChild(opt);
                }
            });

            alCambiarCentro();
        }

        function alCambiarCentro() {
            const selectCentro = document.getElementById('filtro-centro');
            const selectArea = document.getElementById('filtro-area');
            const cenVal = (selectCentro.value || '').trim();

            selectArea.innerHTML = '<option value="">' + (cenVal ? 'Todas las Áreas de ' + cenVal : 'Todas las Áreas') + '</option>';

            let areasDisponibles = [];
            if (cenVal) {
                if (mapaAreasPorCentro[cenVal]) {
                    areasDisponibles = mapaAreasPorCentro[cenVal];
                } else {
                    const key = Object.keys(mapaAreasPorCentro).find(k => k.toLowerCase() === cenVal.toLowerCase());
                    areasDisponibles = key ? mapaAreasPorCentro[key] : [];
                }
            } else {
                areasDisponibles = todasLasAreasOriginales;
            }

            areasDisponibles.forEach(a => {
                const opt = document.createElement('option');
                opt.value = a;
                opt.textContent = a;
                selectArea.appendChild(opt);
            });

            filtrar();
        }

        function filtrar() {
            const redSel = document.getElementById('filtro-red').value.toLowerCase().trim();
            const centroSel = document.getElementById('filtro-centro').value.toLowerCase().trim();
            const areaSel = document.getElementById('filtro-area').value.toLowerCase().trim();
            const busqueda = document.getElementById('busqueda').value.toLowerCase().trim();

            const areaSections = document.querySelectorAll('.area-section');
            let totalVisibles = 0;
            let operativosVisibles = 0;
            let garantiaVisibles = 0;
            let bajasVisibles = 0;
            let countPillTodo = 0;
            let countPillEq = 0;
            let countPillMu = 0;

            areaSections.forEach(section => {
                const areaNombre = section.getAttribute('data-area-name').toLowerCase().trim();
                const tarjetas = section.querySelectorAll('.equipment-card');
                let tarjetasVisiblesEnArea = 0;

                tarjetas.forEach(t => {
                    const tTipo = t.getAttribute('data-tipo') || 'EQUIPO';
                    const tRed = (t.getAttribute('data-red') || '').toLowerCase();
                    const tCentro = (t.getAttribute('data-centro') || '').toLowerCase();
                    const tArea = (t.getAttribute('data-area') || '').toLowerCase();
                    const tEstado = t.getAttribute('data-estado') || '';
                    const tGarantia = t.getAttribute('data-garantia') || '';
                    const tTexto = (t.getAttribute('data-texto') || '').toLowerCase();

                    const matchRed = !redSel || tRed.includes(redSel);
                    const matchCentro = !centroSel || tCentro.includes(centroSel);
                    const matchArea = !areaSel || tArea.includes(areaSel);
                    const matchBusqueda = !busqueda || tTexto.includes(busqueda);

                    if (matchRed && matchCentro && matchArea && matchBusqueda) {
                        countPillTodo++;
                        if (tTipo === 'EQUIPO') countPillEq++;
                        if (tTipo === 'MUEBLE') countPillMu++;
                    }

                    const matchTipo = (TIPO_ACTIVO_FILTRO === 'TODO') || (tTipo === TIPO_ACTIVO_FILTRO);

                    if (matchTipo && matchRed && matchCentro && matchArea && matchBusqueda) {
                        t.style.display = 'block';
                        tarjetasVisiblesEnArea++;
                        totalVisibles++;
                        const estLow = tEstado.toLowerCase();
                        if (estLow.includes('baja') || estLow.includes('mal')) {
                            bajasVisibles++;
                        } else if (estLow.includes('reg') || estLow.includes('man') || tGarantia === 'Con Garantía') {
                            garantiaVisibles++;
                        } else {
                            operativosVisibles++;
                        }
                    } else {
                        t.style.display = 'none';
                    }
                });

                const countLabel = section.querySelector('.count-label');
                if (countLabel) {
                    let txt = `${tarjetasVisiblesEnArea} activo${tarjetasVisiblesEnArea === 1 ? '' : 's'}`;
                    if (TIPO_ACTIVO_FILTRO === 'EQUIPO') {
                        txt = `${tarjetasVisiblesEnArea} equipo${tarjetasVisiblesEnArea === 1 ? '' : 's'} médico${tarjetasVisiblesEnArea === 1 ? '' : 's'}`;
                    } else if (TIPO_ACTIVO_FILTRO === 'MUEBLE') {
                        txt = `${tarjetasVisiblesEnArea} activo${tarjetasVisiblesEnArea === 1 ? '' : 's'} fijo${tarjetasVisiblesEnArea === 1 ? '' : 's'}`;
                    }
                    countLabel.textContent = txt;
                }

                if (tarjetasVisiblesEnArea > 0) {
                    section.style.display = 'block';
                } else {
                    section.style.display = 'none';
                }
            });

            const elPillTodo = document.getElementById('cnt-pills-todo');
            const elPillEq = document.getElementById('cnt-pills-equipos');
            const elPillMu = document.getElementById('cnt-pills-muebles');
            if (elPillTodo) elPillTodo.textContent = countPillTodo;
            if (elPillEq) elPillEq.textContent = countPillEq;
            if (elPillMu) elPillMu.textContent = countPillMu;

            // Actualizar contadores
            document.getElementById('stat-total').textContent = totalVisibles;
            document.getElementById('stat-op').textContent = operativosVisibles;
            document.getElementById('stat-gar').textContent = garantiaVisibles;
            document.getElementById('stat-baj').textContent = bajasVisibles;

            const noRes = document.getElementById('no-results');
            if (totalVisibles === 0) {
                noRes.style.display = 'block';
            } else {
                noRes.style.display = 'none';
            }
        }

        // Ejecutar conteo inicial y pre-seleccionar centro activo si está disponible
        document.addEventListener('DOMContentLoaded', () => {
            const selectCentro = document.getElementById('filtro-centro');
            const urlParams = new URLSearchParams(window.location.search);
            const centroParam = urlParams.get('centro');
            if (centroParam && selectCentro.querySelector(`option[value="${centroParam}"]`)) {
                selectCentro.value = centroParam;
            } else if (selectCentro.querySelector('option[value="BAJO SAN PEDRO"]')) {
                selectCentro.value = "BAJO SAN PEDRO";
            }
            alCambiarCentro();
        });
    </script>
</body>
</html>
"""

def obtener_activos_unificados_db(red_filtro=None, centro_filtro=None):
    """
    Carga de forma unificada los Equipos Médicos y Muebles/Computadoras de la base oficial (2026).
    Normaliza campos territoriales (Red, Centro, Área) y atributos comunes para vistas web y censo.
    """
    from database import calcular_proximos_mantenimientos
    conn = obtener_conexion()
    if not conn:
        return [], [], [], [], [], []
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        
        # 1. Catálogo oficial de redes y centros
        cur.execute("SELECT id, nombre, codigo FROM redes_salud ORDER BY id ASC")
        redes_db = [dict(r) for r in cur.fetchall()]
        redes_dict = {r['id']: r['nombre'] for r in redes_db}
        
        cur.execute("""
            SELECT c.id, c.nombre, c.red_salud_id, r.nombre as red_nombre 
            FROM centros_salud c
            LEFT JOIN redes_salud r ON c.red_salud_id = r.id
            ORDER BY c.nombre ASC
        """)
        centros_db = [dict(r) for r in cur.fetchall()]

        # 2. Cargar Equipos
        cur.execute("SELECT * FROM equipos WHERE COALESCE(estado, 'Operativo') NOT IN ('Inactivo', 'Eliminado') ORDER BY nombre ASC")
        equipos_raw = [dict(r) for r in cur.fetchall()]
        
        # 3. Cargar Muebles y Computación
        cur.execute("""
            SELECT m.*, r.nombre as red_nombre_fk, c.nombre as centro_nombre_fk 
            FROM muebleria m 
            LEFT JOIN redes_salud r ON m.red_salud_id = r.id 
            LEFT JOIN centros_salud c ON m.centro_salud_id = c.id
            WHERE COALESCE(m.estado, 'Activo') NOT IN ('Inactivo', 'Eliminado')
            ORDER BY m.id DESC
        """)
        muebles_raw = [dict(r) for r in cur.fetchall()]
        cur.close()
        conn.close()

        hoy = date.today()
        todos_equipos = []
        for eq in equipos_raw:
            a_nom = (eq.get('area') or eq.get('servicio') or 'General').strip()
            if not a_nom:
                a_nom = 'General'
            eq['area'] = a_nom
            eq['servicio'] = eq.get('servicio') or a_nom
            eq['_tipo'] = 'EQUIPO'
            eq['_tipo_label'] = 'Equipo Médico'
            eq['_icono'] = '🩺'
            eq['id_db'] = eq['id']
            eq['codigo_af'] = str(eq['id'])
            
            # Próximo mantenimiento
            if eq.get('estado') != 'Baja':
                proximos = calcular_proximos_mantenimientos(eq, cantidad=1, hoy=hoy)
                if proximos:
                    eq['f_prox'] = proximos[0].strftime("%d/%m/%Y")
            todos_equipos.append(eq)

        # Identificadores de equipos médicos para evitar duplicados en muebles
        EXENTOS_DUPLICADOS = {"", "S/C", "0", "DONACION", "SIN CODIGO", "SIN SERIE", "NINGUNO", "N/A", "NO APLICA", "S/N", "SN", "-"}
        series_eq = {str(eq.get('numero_serie') or '').strip().upper() for eq in equipos_raw if str(eq.get('numero_serie') or '').strip().upper() not in EXENTOS_DUPLICADOS}
        sispam_eq = {str(eq.get('codigo_sispam') or '').strip().upper() for eq in equipos_raw if str(eq.get('codigo_sispam') or '').strip().upper() not in EXENTOS_DUPLICADOS}
        bertin_eq = {str(eq.get('bertin') or '').strip().upper() for eq in equipos_raw if str(eq.get('bertin') or '').strip().upper() not in EXENTOS_DUPLICADOS}
        sapm_eq = {str(eq.get('sapm') or '').strip().upper() for eq in equipos_raw if str(eq.get('sapm') or '').strip().upper() not in EXENTOS_DUPLICADOS}

        todos_muebles = []
        for m in muebles_raw:
            m_serie = str(m.get('serie') or '').strip().upper()
            m_sispam = str(m.get('codigo_sispam') or '').strip().upper()
            m_bertin = str(m.get('bertin') or '').strip().upper()
            m_sapm = str(m.get('sapm') or '').strip().upper()

            # Evitar mostrar muebles que en realidad son clones duplicados de un equipo médico
            if (m_serie and m_serie in series_eq) or \
               (m_sispam and m_sispam in sispam_eq) or \
               (m_bertin and m_bertin in bertin_eq) or \
               (m_sapm and m_sapm in sapm_eq):
                continue

            # Resolver Centro de Salud
            cen_nom = m.get('centro_nombre_fk') or m.get('unidad_organizacional') or ''
            if not cen_nom:
                ub = m.get('ubicacion') or ''
                for c in centros_db:
                    if c['nombre'].upper() in ub.upper():
                        cen_nom = c['nombre']
                        break
            if not cen_nom:
                cen_nom = 'Centro de Salud'

            # Resolver Red de Salud
            red_nom = m.get('red_nombre_fk') or m.get('direccion_administrativa') or ''
            if not red_nom:
                for c in centros_db:
                    if c['nombre'].upper() == cen_nom.upper():
                        red_nom = c.get('red_nombre') or redes_dict.get(c['red_salud_id'], '')
                        break
            if not red_nom:
                red_nom = 'Red GAMLP'

            # Resolver Área: Mantener el nombre completo tal como está registrado (incluyendo piso ej: Piso 2 - Ejemplo Multifuncional)
            ub = (m.get('ubicacion') or 'General').strip()
            area_nom = ub if ub else 'General'

            # Icono según descripción o tipo
            desc_lower = (str(m.get('descripcion') or '') + ' ' + str(m.get('tipo_activo') or '')).lower()
            if any(w in desc_lower for w in ['compu', 'monitor', 'laptop', 'pc', 'teclado', 'cpu', 'servidor', 'impresora']):
                icono = '💻'
            elif any(w in desc_lower for w in ['desfib', 'desfrib', 'electro', 'aspirad', 'monitor de signos', 'tensio', 'oxim']):
                icono = '🩺'
            else:
                icono = '🛋️'

            cod_af = m.get('codigo_sispam')
            if not cod_af or cod_af == 'S/C':
                cod_af = m.get('bertin') or m.get('sapm') or f"MUE-{m['id']:03d}"

            item_mueble = {
                'id': f"MUE-{m['id']}",
                'id_db': m['id'],
                'codigo_af': cod_af,
                'nombre': m.get('descripcion') or m.get('tipo_activo') or 'Mueble / Computación',
                'tipo_activo': m.get('tipo_activo') or 'Mueble / TI',
                'marca': m.get('marca') or 'S/M',
                'modelo': m.get('modelo') or 'S/M',
                'numero_serie': m.get('serie') or 'S/N',
                'area': area_nom,
                'servicio': area_nom,
                'red_salud_nombre': red_nom,
                'centro_salud_nombre': cen_nom,
                'estado': m.get('estado_conservacion') or m.get('estado') or 'Bueno',
                'garantia': '',
                'responsable': m.get('persona_asignada') or '',
                'cargo': m.get('cargo_asignado') or '',
                'ci': m.get('ci_asignado') or '',
                'tecnico': m.get('tecnico_inventareador') or '',
                'fecha_asignacion': m.get('fecha_asignacion') or '',
                'observaciones': m.get('observaciones_de_asignacion') or '',
                'codigo_sispam': m.get('codigo_sispam') or '',
                'bertin': m.get('bertin') or '',
                'sapm': m.get('sapm') or '',
                'foto': None,
                '_tipo': 'MUEBLE',
                '_tipo_label': 'Mueble / TI',
                '_icono': icono
            }
            todos_muebles.append(item_mueble)

        # Unificar nombres de áreas para que coincidan exactamente (con su piso) en el mismo centro
        todas_las_areas = set()
        for eq in todos_equipos:
            if eq.get('area'): todas_las_areas.add(eq['area'].strip())
        for m in todos_muebles:
            if m.get('area'): todas_las_areas.add(m['area'].strip())

        def _unificar_nombre_area(ar_str):
            if not ar_str or ar_str.lower() in ('general', '-', 'ninguno'):
                return 'General'
            ar_clean = ar_str.strip()
            if any(ar_clean.lower().startswith(p) for p in ['piso ', 'pb ', 'pb-', 'nivel ', 'planta ']):
                return ar_clean
            for a_full in todas_las_areas:
                if ' - ' in a_full and a_full.split(' - ', 1)[1].strip().lower() == ar_clean.lower():
                    return a_full
            return ar_clean

        for eq in todos_equipos:
            eq['area'] = _unificar_nombre_area(eq.get('area'))
        for m in todos_muebles:
            m['area'] = _unificar_nombre_area(m.get('area'))

        activos_unificados = todos_equipos + todos_muebles

        # Filtrar si se solicitaron filtros territoriales
        if red_filtro:
            r_low = red_filtro.lower()
            activos_unificados = [a for a in activos_unificados if r_low in str(a.get('red_salud_nombre', '')).lower()]
            todos_equipos = [a for a in todos_equipos if r_low in str(a.get('red_salud_nombre', '')).lower()]
            todos_muebles = [a for a in todos_muebles if r_low in str(a.get('red_salud_nombre', '')).lower()]

        if centro_filtro:
            c_low = centro_filtro.lower()
            activos_unificados = [a for a in activos_unificados if c_low in str(a.get('centro_salud_nombre', '')).lower()]
            todos_equipos = [a for a in todos_equipos if c_low in str(a.get('centro_salud_nombre', '')).lower()]
            todos_muebles = [a for a in todos_muebles if c_low in str(a.get('centro_salud_nombre', '')).lower()]

        # Áreas presentes
        areas_set = set()
        for a in activos_unificados:
            ar = a.get('area')
            if ar and ar.strip():
                areas_set.add(ar.strip())
        areas_lista = sorted(list(areas_set))
        if not areas_lista:
            areas_lista = ['General']

        return activos_unificados, todos_equipos, todos_muebles, redes_db, centros_db, areas_lista
    except Exception as e:
        print(f"[ERROR] obtener_activos_unificados_db: {e}")
        try: conn.close()
        except: pass
        return [], [], [], [], [], []

@app_web.route('/')
@app_web.route('/inventario')
def vista_inventario_web():
    try:
        activos_db, eqs_db, mus_db, redes_db, centros_db, areas_lista = obtener_activos_unificados_db()

        total = len(activos_db)
        cnt_todo = len(activos_db)
        cnt_equipos = len(eqs_db)
        cnt_muebles = len(mus_db)

        operativos = sum(1 for a in activos_db if 'baja' not in str(a.get('estado','')).lower() and 'mal' not in str(a.get('estado','')).lower())
        garantia = sum(1 for a in activos_db if a.get('garantia') == 'Con Garantía' or 'reg' in str(a.get('estado','')).lower() or 'man' in str(a.get('estado','')).lower())
        bajas = sum(1 for a in activos_db if 'baja' in str(a.get('estado','')).lower() or 'mal' in str(a.get('estado','')).lower())

        # Mapear qué áreas pertenecen a qué centro de salud
        mapa_areas_centro = {}
        for a in activos_db:
            cen = (a.get('centro_salud_nombre') or '').strip()
            ar = (a.get('area') or '').strip()
            if cen and ar:
                mapa_areas_centro.setdefault(cen, set()).add(ar)

        areas_bd = obtener_areas_db()
        for ab in areas_bd:
            cen = (ab.get('centro_salud_nombre') or '').strip()
            nom = (ab.get('nombre') or '').strip()
            piso = (ab.get('piso') or '').strip()
            if not nom: continue
            if piso and piso != '-':
                p_low = piso.lower()
                if not p_low.startswith('piso') and not p_low.startswith('planta') and not p_low.startswith('pb'):
                    piso = f"Piso {piso}"
                fmt_ar = f"{piso} - {nom}"
            else:
                fmt_ar = nom
            if cen and fmt_ar:
                mapa_areas_centro.setdefault(cen, set()).add(fmt_ar)

        mapa_areas_json = {k: sorted(list(v)) for k, v in mapa_areas_centro.items()}

        return render_template_string(
            HTML_INVENTARIO, 
            activos=activos_db,
            equipos=eqs_db,
            muebles=mus_db,
            redes=redes_db,
            centros=centros_db,
            areas=areas_lista,
            mapa_areas_centro=json.dumps(mapa_areas_json),
            total=total, 
            cnt_todo=cnt_todo,
            cnt_equipos=cnt_equipos,
            cnt_muebles=cnt_muebles,
            operativos=operativos, 
            garantia=garantia, 
            bajas=bajas
        )
    except Exception as e:
        return f"Error cargando inventario: {e}", 500

HTML_ANALISIS = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>SGEM GAMLP - Análisis y Censo de Equipamiento</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        :root {
            --primary: #0F172A;
            --primary-dark: #020617;
            --primary-hover: #1E293B;
            --bg: #F8FAFC;
            --card-bg: #FFFFFF;
            --text: #0F172A;
            --muted: #64748B;
            --border: #E2E8F0;
            --border-hover: #CBD5E1;
            --green: #10B981;
            --orange: #F59E0B;
            --purple: #8B5CF6;
            --red: #EF4444;
        }

        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background-color: var(--bg);
            color: var(--text);
            margin: 0;
            padding: 0;
            -webkit-font-smoothing: antialiased;
        }

        .header {
            background: #0F172A;
            color: white;
            padding: 24px 20px 20px;
            text-align: center;
            border-bottom: 1px solid #1E293B;
            box-shadow: 0 1px 3px rgba(0,0,0,0.1);
        }

        .badge-gamlp {
            display: inline-block;
            background: rgba(255, 255, 255, 0.08);
            border: 1px solid rgba(255, 255, 255, 0.14);
            color: #CBD5E1;
            font-size: 11px;
            font-weight: 700;
            padding: 3px 12px;
            border-radius: 20px;
            margin-bottom: 8px;
            letter-spacing: 0.5px;
            text-transform: uppercase;
        }

        .header h1 {
            font-size: 22px;
            font-weight: 800;
            margin: 0 0 4px;
            letter-spacing: -0.5px;
        }

        .header p {
            font-size: 13px;
            color: #94A3B8;
            margin: 0;
        }

        .nav-tabs {
            display: flex;
            justify-content: center;
            gap: 8px;
            margin-top: 15px;
        }

        .nav-tab {
            color: #94A3B8;
            text-decoration: none;
            padding: 6px 14px;
            border-radius: 20px;
            font-size: 12.5px;
            font-weight: 600;
            background: rgba(255, 255, 255, 0.06);
            border: 1px solid rgba(255, 255, 255, 0.1);
            transition: all 0.2s;
        }

        .nav-tab:hover {
            color: #FFFFFF;
            border-color: rgba(255, 255, 255, 0.25);
        }

        .nav-tab.active {
            background: #FFFFFF;
            color: #0F172A;
            font-weight: 700;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.2);
        }

        .asset-tab {
            padding: 7px 16px;
            border-radius: 20px;
            font-size: 13px;
            font-weight: 700;
            text-decoration: none;
            color: #475569;
            background: #FFFFFF;
            border: 1.5px solid #E2E8F0;
            transition: all 0.2s;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }
        .asset-tab:hover {
            border-color: #0F172A;
            color: #0F172A;
        }
        .asset-tab.active {
            background: #0F172A;
            color: #FFFFFF;
            border-color: #0F172A;
            box-shadow: 0 2px 6px rgba(15, 23, 42, 0.25);
        }

        .container {
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px 16px 60px;
        }

        .filter-card {
            background: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 16px;
            margin-bottom: 20px;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.03);
        }

        .filter-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
            gap: 12px;
        }

        .filter-group {
            display: flex;
            flex-direction: column;
            gap: 6px;
        }

        .filter-label {
            font-size: 12px;
            font-weight: 700;
            color: var(--muted);
            text-transform: uppercase;
        }

        .select-input {
            width: 100%;
            padding: 10px 12px;
            border: 1.5px solid var(--border);
            border-radius: 10px;
            font-size: 14px;
            font-family: inherit;
            background: white;
            color: var(--text);
            box-sizing: border-box;
            outline: none;
            font-weight: 600;
        }

        .stats-summary {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
            gap: 12px;
            margin-bottom: 24px;
        }

        .stat-box {
            background: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 14px;
            padding: 16px;
            text-align: center;
            box-shadow: 0 2px 6px rgba(0,0,0,0.02);
        }

        .stat-box .val {
            font-size: 26px;
            font-weight: 800;
            color: var(--primary);
        }

        .stat-box .lbl {
            font-size: 12px;
            font-weight: 700;
            color: var(--muted);
            margin-top: 4px;
            text-transform: uppercase;
        }

        .section-header {
            margin-bottom: 14px;
        }

        .section-header h2 {
            font-size: 18px;
            font-weight: 800;
            margin: 0;
            color: var(--text);
        }

        .section-header p {
            font-size: 13px;
            color: var(--muted);
            margin: 4px 0 0;
        }

        .cards-grid {
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(260px, 1fr));
            gap: 14px;
            margin-bottom: 24px;
        }

        .censo-card {
            background: white;
            border: 1px solid var(--border);
            border-radius: 14px;
            padding: 16px;
            cursor: pointer;
            transition: transform 0.15s, box-shadow 0.15s, border-color 0.15s;
            position: relative;
            overflow: hidden;
            display: flex;
            flex-direction: column;
            justify-content: space-between;
        }

        .censo-card:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 20px rgba(0, 0, 0, 0.06);
            border-color: var(--primary);
        }

        .censo-card .top-bar {
            position: absolute;
            top: 0;
            left: 0;
            right: 0;
            height: 4px;
            background: var(--primary);
        }

        .censo-card .card-title {
            font-size: 15px;
            font-weight: 700;
            color: var(--text);
            margin-bottom: 8px;
        }

        .censo-card .card-count {
            font-size: 24px;
            font-weight: 800;
            color: var(--primary);
            margin-bottom: 12px;
        }

        .censo-card .btn-inspect {
            font-size: 12px;
            font-weight: 700;
            color: white;
            background: var(--primary);
            padding: 8px 12px;
            border-radius: 8px;
            text-align: center;
            text-decoration: none;
            display: block;
        }

        .charts-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(360px, 1fr));
            gap: 16px;
            margin-bottom: 30px;
        }

        .chart-box {
            background: white;
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 18px;
            box-shadow: 0 2px 8px rgba(0,0,0,0.02);
        }

        .chart-box h3 {
            font-size: 15px;
            font-weight: 700;
            margin: 0 0 14px;
            color: var(--text);
        }

        /* Modal */
        .modal-overlay {
            display: none;
            position: fixed;
            top: 0; left: 0; right: 0; bottom: 0;
            background: rgba(15, 23, 42, 0.6);
            z-index: 9999;
            align-items: center;
            justify-content: center;
            padding: 16px;
            backdrop-filter: blur(4px);
        }

        .modal-content {
            background: white;
            border-radius: 16px;
            max-width: 900px;
            width: 100%;
            max-height: 85vh;
            display: flex;
            flex-direction: column;
            box-shadow: 0 20px 40px rgba(0,0,0,0.2);
            overflow: hidden;
        }

        .modal-header {
            padding: 16px 20px;
            border-bottom: 1px solid var(--border);
            display: flex;
            justify-content: space-between;
            align-items: center;
            background: #F8FAFC;
        }

        .modal-header h3 {
            margin: 0;
            font-size: 17px;
            font-weight: 800;
        }

        .modal-body {
            padding: 16px 20px;
            overflow-y: auto;
            flex: 1;
        }

        .modal-search {
            width: 100%;
            padding: 10px 14px;
            border: 1.5px solid var(--border);
            border-radius: 10px;
            font-size: 14px;
            margin-bottom: 14px;
            box-sizing: border-box;
            outline: none;
        }

        .table-responsive {
            width: 100%;
            border-collapse: collapse;
            font-size: 13px;
        }

        .table-responsive th {
            background: #F1F5F9;
            padding: 10px 12px;
            text-align: left;
            font-weight: 700;
            color: var(--muted);
            border-bottom: 1.5px solid var(--border);
        }

        .table-responsive td {
            padding: 10px 12px;
            border-bottom: 1px solid var(--border);
            color: var(--text);
        }

        .table-responsive tr:hover {
            background: #F8FAFC;
        }

        .btn-view-link {
            display: inline-block;
            background: #EFF6FF;
            color: var(--primary);
            padding: 4px 8px;
            border-radius: 6px;
            font-size: 11px;
            font-weight: 700;
            text-decoration: none;
        }

        .btn-close-modal {
            background: #E2E8F0;
            border: none;
            border-radius: 8px;
            padding: 8px 14px;
            font-weight: 700;
            cursor: pointer;
            font-size: 13px;
        }
    </style>
</head>
<body>
    <div class="header">
        <span class="badge-gamlp">GAMLP • SGEM v1.1</span>
        <h1>Sistema de Gestión de Equipamiento Médico</h1>
        <p>Análisis Estadístico y Censo Territorial</p>
        <div class="nav-tabs">
            <a href="/inventario" class="nav-tab">📦 Inventario</a>
            <a href="/analisis" class="nav-tab active">📊 Análisis y Censo</a>
            <a href="/movil" class="nav-tab">📱 Registro Móvil</a>
            <a href="/historico" class="nav-tab">🏛️ Histórico</a>
            <a href="/enlaces" class="nav-tab">🔗 Enlaces</a>
        </div>
    </div>

    <div class="container">
        <!-- Selector Tipo de Activo en Análisis -->
        <div style="display: flex; gap: 10px; margin-bottom: 18px; flex-wrap: wrap;">
            <a href="javascript:void(0)" class="asset-tab {% if tipo_sel == 'TODO' %}active{% endif %}" onclick="cambiarTipoActivo('TODO')">
                🌐 Ver Todo ({{ cnt_todo }})
            </a>
            <a href="javascript:void(0)" class="asset-tab {% if tipo_sel == 'EQUIPOS' %}active{% endif %}" onclick="cambiarTipoActivo('EQUIPOS')">
                🩺 Solo Equipos Médicos ({{ cnt_equipos }})
            </a>
            <a href="javascript:void(0)" class="asset-tab {% if tipo_sel == 'MUEBLES' %}active{% endif %}" onclick="cambiarTipoActivo('MUEBLES')">
                🛋️ Solo Muebles y TI ({{ cnt_muebles }})
            </a>
        </div>

        <!-- Filtros Territoriales -->
        <div class="filter-card">
            <div class="filter-grid">
                <div class="filter-group">
                    <label class="filter-label">🌐 Red de Salud</label>
                    <select id="filtro-red" class="select-input" onchange="alCambiarRedWeb()">
                        <option value="">Todas las Redes (GAMLP)</option>
                        {% for r in redes %}
                        <option value="{{ r['nombre'] }}" data-id="{{ r['id'] }}" {% if red_sel == r['nombre'] %}selected{% endif %}>{{ r['nombre'] }}</option>
                        {% endfor %}
                    </select>
                </div>
                <div class="filter-group">
                    <label class="filter-label">🏥 Centro de Salud</label>
                    <select id="filtro-centro" class="select-input" onchange="alCambiarCentroWeb()">
                        <option value="">Todos los Centros</option>
                        {% for c in centros %}
                        <option value="{{ c['nombre'] }}" data-red-id="{{ c['red_salud_id'] }}" data-red-name="{{ c['red_nombre'] }}" {% if centro_sel == c['nombre'] %}selected{% endif %}>{{ c['nombre'] }}</option>
                        {% endfor %}
                    </select>
                </div>
            </div>
        </div>

        <!-- Resumen de Estadísticas -->
        <div class="stats-summary">
            <div class="stat-box">
                <div class="val" id="stat-total">{{ total }}</div>
                <div class="lbl">Total Activos</div>
            </div>
            <div class="stat-box">
                <div class="val" style="color: var(--green);" id="stat-op">{{ operativos }}</div>
                <div class="lbl">Operativos / Buen Estado</div>
            </div>
            <div class="stat-box">
                <div class="val" style="color: var(--purple);" id="stat-gar">{{ garantia }}</div>
                <div class="lbl">En Garantía / Regular</div>
            </div>
            <div class="stat-box">
                <div class="val" style="color: var(--red);" id="stat-baj">{{ bajas }}</div>
                <div class="lbl">Bajas / Fuera de Servicio</div>
            </div>
        </div>

        <!-- Sección de Censo Jerárquico Dinámico -->
        <div class="section-header">
            <h2 id="censo-titulo">{{ censo_titulo }}</h2>
            <p id="censo-subtitulo">{{ censo_subtitulo }}</p>
        </div>

        <div class="cards-grid" id="censo-grid">
            {% for item in censo_items %}
            <div class="censo-card" onclick="inspeccionarGrupo('{{ item.nombre | escape }}')">
                <div class="top-bar"></div>
                <div class="card-title">{{ item.nombre_display }}</div>
                <div class="card-count">{{ item.cantidad }} <span style="font-size: 13px; font-weight: 600; color: var(--muted);">activos</span></div>
                <a class="btn-inspect" href="javascript:void(0)">🔍 Ver {{ item.tipo_label }}</a>
            </div>
            {% endfor %}
        </div>

        <!-- Gráficas Visuales Interactivas (Chart.js) -->
        <div class="charts-grid">
            <div class="chart-box">
                <h3>📊 {{ censo_titulo }}</h3>
                <div style="position: relative; height: 260px;">
                    <canvas id="chart-censo"></canvas>
                </div>
            </div>
            <div class="chart-box">
                <h3>📋 Tipos de Activos más Frecuentes</h3>
                <div style="position: relative; height: 260px;">
                    <canvas id="chart-tipos"></canvas>
                </div>
            </div>
        </div>

        <!-- Tabla de Inventario de Activos en la misma página -->
        <div style="background: white; border: 1px solid var(--border); border-radius: 16px; padding: 18px; box-shadow: 0 2px 8px rgba(0,0,0,0.02); margin-bottom: 30px;">
            <div style="display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 10px; margin-bottom: 14px;">
                <div>
                    <h3 style="font-size: 16px; font-weight: 800; margin: 0; color: var(--text);">📋 Inventario y Listado de Activos</h3>
                    <p style="font-size: 13px; color: var(--muted); margin: 3px 0 0;">{{ total }} activos encontrados en este filtro territorial</p>
                </div>
                <input type="text" id="inline-buscar" class="select-input" style="max-width: 280px; padding: 8px 12px; font-size: 13px;" placeholder="🔍 Filtrar en esta tabla..." oninput="filtrarTablaInline()">
            </div>
            <div style="overflow-x: auto; max-height: 440px;">
                <table class="table-responsive">
                    <thead>
                        <tr>
                            <th>Tipo</th>
                            <th>Cod. AF / SISFAM</th>
                            <th>Descripción / Nombre</th>
                            <th>Marca</th>
                            <th>Modelo</th>
                            <th>Centro de Salud</th>
                            <th>Área / Ubicación</th>
                            <th>Estado</th>
                            <th>Ficha</th>
                        </tr>
                    </thead>
                    <tbody id="inline-tbody">
                        {% for eq in eqs_vista %}
                        <tr>
                            <td>
                                <span style="font-size: 16px;">{{ eq.get('_icono', '🩺') }}</span>
                                <span style="font-size: 11px; font-weight: 700; color: var(--muted);">{{ eq.get('_tipo_label', 'Activo') }}</span>
                            </td>
                            <td><strong>{{ eq.get('codigo_af') or eq['id'] }}</strong></td>
                            <td>{{ eq['nombre'] }}</td>
                            <td>{{ eq.get('marca') or '-' }}</td>
                            <td>{{ eq.get('modelo') or '-' }}</td>
                            <td>{{ eq.get('centro_salud_nombre') or '-' }}</td>
                            <td>{{ eq.get('area') or eq.get('servicio') or 'General' }}</td>
                            <td>
                                {% set est = (eq.get('estado') or 'Bueno')|lower %}
                                {% if 'baja' in est or 'mal' in est %}
                                <span style="background: #FEE2E2; color: #991B1B; padding: 3px 8px; border-radius: 12px; font-weight: 700; font-size: 11px;">Baja</span>
                                {% elif 'reg' in est or 'man' in est %}
                                <span style="background: #FEF3C7; color: #92400E; padding: 3px 8px; border-radius: 12px; font-weight: 700; font-size: 11px;">Regular</span>
                                {% else %}
                                <span style="background: #ECFDF5; color: #065F46; padding: 3px 8px; border-radius: 12px; font-weight: 700; font-size: 11px;">Operativo</span>
                                {% endif %}
                            </td>
                            <td>
                                {% if eq.get('_tipo') == 'MUEBLE' %}
                                <a href="/mueble/{{ eq.get('id_db', eq['id']) }}" class="btn-view-link" target="_blank" style="background: #FFFBEB; color: #B45309; border: 1px solid #FDE68A;">🛋️ Ficha Mueble</a>
                                {% else %}
                                <a href="/equipo/{{ eq['id'] }}" class="btn-view-link" target="_blank">🩺 Ficha Equipo</a>
                                {% endif %}
                            </td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>

    <!-- Modal de Detalle de Equipos -->
    <div class="modal-overlay" id="modal-detalle">
        <div class="modal-content">
            <div class="modal-header">
                <h3 id="modal-titulo">Equipos en: </h3>
                <button class="btn-close-modal" onclick="cerrarModal()">✕ Cerrar</button>
            </div>
            <div class="modal-body">
                <input type="text" id="modal-buscar" class="modal-search" placeholder="🔍 Buscar por nombre, marca, modelo, Cod. AF..." oninput="filtrarModal()">
                <div style="overflow-x: auto;">
                    <table class="table-responsive">
                        <thead>
                            <tr>
                                <th>Cod. AF</th>
                                <th>Equipo Médico</th>
                                <th>Marca</th>
                                <th>Modelo</th>
                                <th>Centro de Salud</th>
                                <th>Área/Servicio</th>
                                <th>Acción</th>
                            </tr>
                        </thead>
                        <tbody id="modal-tbody">
                        </tbody>
                    </table>
                </div>
            </div>
        </div>
    </div>

    <script>
        const todosActivos = {{ activos_json | safe }};
        let activosModalActual = [];

        function cambiarTipoActivo(tipo) {
            const redSel = document.getElementById('filtro-red').value;
            const centroSel = document.getElementById('filtro-centro').value;
            window.location.href = `/analisis?tipo=${encodeURIComponent(tipo)}&red=${encodeURIComponent(redSel)}&centro=${encodeURIComponent(centroSel)}`;
        }

        function filtrarTablaInline() {
            const query = document.getElementById('inline-buscar').value.toLowerCase().trim();
            const rows = document.querySelectorAll('#inline-tbody tr');
            rows.forEach(tr => {
                const text = tr.innerText.toLowerCase();
                tr.style.display = (!query || text.includes(query)) ? '' : 'none';
            });
        }

        function alCambiarRedWeb() {
            const redSel = document.getElementById('filtro-red').value;
            const centroSelect = document.getElementById('filtro-centro');
            const tipoActual = '{{ tipo_sel }}';
            
            const redOption = document.getElementById('filtro-red').selectedOptions[0];
            const redId = redOption ? redOption.getAttribute('data-id') : '';

            Array.from(centroSelect.options).forEach(opt => {
                if (opt.value === '') return;
                const optRedId = opt.getAttribute('data-red-id');
                if (!redId || optRedId === redId) {
                    opt.style.display = 'block';
                } else {
                    opt.style.display = 'none';
                }
            });

            const selectedCentroOpt = centroSelect.selectedOptions[0];
            if (selectedCentroOpt && selectedCentroOpt.style.display === 'none') {
                centroSelect.value = '';
            }

            const centroVal = centroSelect.value;
            window.location.href = `/analisis?tipo=${encodeURIComponent(tipoActual)}&red=${encodeURIComponent(redSel)}&centro=${encodeURIComponent(centroVal)}`;
        }

        function alCambiarCentroWeb() {
            const redSel = document.getElementById('filtro-red').value;
            const centroSel = document.getElementById('filtro-centro').value;
            const tipoActual = '{{ tipo_sel }}';
            window.location.href = `/analisis?tipo=${encodeURIComponent(tipoActual)}&red=${encodeURIComponent(redSel)}&centro=${encodeURIComponent(centroSel)}`;
        }

        function inspeccionarGrupo(nombreGrupo) {
            const redSel = document.getElementById('filtro-red').value;
            const centroSel = document.getElementById('filtro-centro').value;

            if (!redSel && !centroSel) {
                document.getElementById('filtro-red').value = nombreGrupo;
                alCambiarRedWeb();
            } else if (redSel && !centroSel) {
                document.getElementById('filtro-centro').value = nombreGrupo;
                alCambiarCentroWeb();
            } else {
                abrirModalArea(nombreGrupo);
            }
        }

        function abrirModalArea(nombreArea) {
            const centroSel = document.getElementById('filtro-centro').value;

            activosModalActual = todosActivos.filter(a => {
                const matchCen = !centroSel || (a.centro_salud_nombre && a.centro_salud_nombre.toLowerCase() === centroSel.toLowerCase());
                const aNom = a.area || a.servicio || 'General';
                const matchArea = aNom.toLowerCase() === nombreArea.toLowerCase();
                return matchCen && matchArea;
            });

            document.getElementById('modal-titulo').textContent = `Activos en Área: ${nombreArea} (${activosModalActual.length} activos)`;
            document.getElementById('modal-buscar').value = '';
            renderizarTablaModal(activosModalActual);
            document.getElementById('modal-detalle').style.display = 'flex';
        }

        function renderizarTablaModal(lista) {
            const tbody = document.getElementById('modal-tbody');
            tbody.innerHTML = '';
            if (lista.length === 0) {
                tbody.innerHTML = '<tr><td colspan="7" style="text-align: center; color: var(--muted); padding: 20px;">No se encontraron activos</td></tr>';
                return;
            }
            lista.forEach(a => {
                const tr = document.createElement('tr');
                const icono = a._icono || (a._tipo === 'MUEBLE' ? '🛋️' : '🩺');
                const cod = a.codigo_af || a.id || '-';
                const link = a._tipo === 'MUEBLE' ? `/mueble/${a.id_db || a.id}` : `/equipo/${encodeURIComponent(a.id)}`;
                tr.innerHTML = `
                    <td><span style="font-size: 15px;">${icono}</span> <strong>${cod}</strong></td>
                    <td>${a.nombre || '-'}</td>
                    <td>${a.marca || '-'}</td>
                    <td>${a.modelo || '-'}</td>
                    <td>${a.centro_salud_nombre || '-'}</td>
                    <td>${a.area || a.servicio || 'General'}</td>
                    <td><a href="${link}" class="btn-view-link" target="_blank">📄 Ver Ficha</a></td>
                `;
                tbody.appendChild(tr);
            });
        }

        function filtrarModal() {
            const query = document.getElementById('modal-buscar').value.toLowerCase().trim();
            if (!query) {
                renderizarTablaModal(activosModalActual);
                return;
            }
            const filtrados = activosModalActual.filter(a => {
                return (
                    (a.codigo_af && String(a.codigo_af).toLowerCase().includes(query)) ||
                    (a.id && String(a.id).toLowerCase().includes(query)) ||
                    (a.nombre && a.nombre.toLowerCase().includes(query)) ||
                    (a.marca && a.marca.toLowerCase().includes(query)) ||
                    (a.modelo && a.modelo.toLowerCase().includes(query)) ||
                    (a.numero_serie && a.numero_serie.toLowerCase().includes(query))
                );
            });
            renderizarTablaModal(filtrados);
        }

        function cerrarModal() {
            document.getElementById('modal-detalle').style.display = 'none';
        }

        window.onclick = function(event) {
            const modal = document.getElementById('modal-detalle');
            if (event.target === modal) {
                cerrarModal();
            }
        };

        // Renderizar Gráficas Visuales con Chart.js
        document.addEventListener('DOMContentLoaded', () => {
            const censoLabels = {{ chart_censo_labels | safe }};
            const censoData = {{ chart_censo_data | safe }};

            const ctxCenso = document.getElementById('chart-censo');
            if (ctxCenso && censoLabels.length > 0) {
                new Chart(ctxCenso, {
                    type: 'bar',
                    data: {
                        labels: censoLabels,
                        datasets: [{
                            label: 'Activos Registrados',
                            data: censoData,
                            backgroundColor: ['#0F172A', '#334155', '#475569', '#64748B', '#10B981', '#F59E0B', '#EF4444', '#0284C7'],
                            borderRadius: 8
                        }]
                    },
                    options: {
                        responsive: true,
                        maintainAspectRatio: false,
                        plugins: {
                            legend: { display: false }
                        },
                        scales: {
                            y: { beginAtZero: true, grid: { color: '#F1F5F9' } },
                            x: { grid: { display: false } }
                        }
                    }
                });
            }

            const tiposLabels = {{ chart_tipos_labels | safe }};
            const tiposData = {{ chart_tipos_data | safe }};

            const ctxTipos = document.getElementById('chart-tipos');
            if (ctxTipos && tiposLabels.length > 0) {
                new Chart(ctxTipos, {
                    type: 'bar',
                    data: {
                        labels: tiposLabels,
                        datasets: [{
                            label: 'Cantidad',
                            data: tiposData,
                            backgroundColor: '#334155',
                            borderRadius: 8
                        }]
                    },
                    options: {
                        indexAxis: 'y',
                        responsive: true,
                        maintainAspectRatio: false,
                        plugins: {
                            legend: { display: false }
                        },
                        scales: {
                            x: { beginAtZero: true, grid: { color: '#F1F5F9' } },
                            y: { grid: { display: false } }
                        }
                    }
                });
            }
        });
    </script>
</body>
</html>
"""

def simplificar_red_web(red_str):
    import re
    if not red_str:
        return "Sin Red"
    s = str(red_str).strip()
    m = re.search(r'RED\s*([0-9]+)', s, re.IGNORECASE)
    if m:
        return f"Red {m.group(1)}"
    if "(" in s:
        s = s.split("(")[0].strip()
    return s.title() if s.isupper() else s

@app_web.route('/analisis')
def vista_analisis_web():
    from collections import Counter
    import json
    tipo_param = request.args.get('tipo', 'TODO').strip().upper()
    if tipo_param not in ('TODO', 'EQUIPOS', 'MUEBLES'):
        tipo_param = 'TODO'
    red_param = request.args.get('red', '').strip()
    centro_param = request.args.get('centro', '').strip()

    try:
        activos_unificados, todos_equipos, todos_muebles, redes_db, centros_db, areas_lista = obtener_activos_unificados_db()

        # Filtrar territorialmente
        activos_en_red = activos_unificados
        if red_param:
            activos_en_red = [a for a in activos_unificados if str(a.get('red_salud_nombre', '')).strip().lower() == red_param.lower()]

        activos_en_centro = activos_en_red
        if centro_param:
            activos_en_centro = [a for a in activos_en_red if str(a.get('centro_salud_nombre', '')).strip().lower() == centro_param.lower()]

        # Conteos en el contexto territorial actual
        cnt_todo = len(activos_en_centro)
        cnt_equipos = sum(1 for a in activos_en_centro if a.get('_tipo') == 'EQUIPO')
        cnt_muebles = sum(1 for a in activos_en_centro if a.get('_tipo') == 'MUEBLE')

        # Filtrar por Tipo de Activo
        if tipo_param == 'EQUIPOS':
            activos_contexto = [a for a in activos_en_centro if a.get('_tipo') == 'EQUIPO']
            tipo_label_censo = "Equipos Médicos"
            activos_censo_gamlp = todos_equipos
        elif tipo_param == 'MUEBLES':
            activos_contexto = [a for a in activos_en_centro if a.get('_tipo') == 'MUEBLE']
            tipo_label_censo = "Muebles y Computación"
            activos_censo_gamlp = todos_muebles
        else:
            activos_contexto = list(activos_en_centro)
            tipo_label_censo = "Activos y Equipamiento"
            activos_censo_gamlp = activos_unificados

        total = len(activos_contexto)
        operativos = sum(1 for a in activos_contexto if 'baja' not in str(a.get('estado', '')).lower() and 'mal' not in str(a.get('estado', '')).lower())
        garantia = sum(1 for a in activos_contexto if a.get('garantia') == 'Con Garantía' or 'reg' in str(a.get('estado', '')).lower() or 'man' in str(a.get('estado', '')).lower())
        bajas = sum(1 for a in activos_contexto if 'baja' in str(a.get('estado', '')).lower() or 'mal' in str(a.get('estado', '')).lower())

        # Determinar Censo Jerárquico
        censo_items = []
        if not red_param and not centro_param:
            # Nivel 1: Todas las redes
            censo_titulo = f"🌐 Censo de {tipo_label_censo} por Red de Salud"
            censo_subtitulo = f"Total en GAMLP: {total:,} registros | Haz clic en una Red para explorar sus Centros de Salud"
            grupos = {}
            for a in activos_censo_gamlp:
                r_nom = a.get('red_salud_nombre') or 'Sin Red'
                grupos.setdefault(r_nom, []).append(a)
            for r_nom, lista in sorted(grupos.items(), key=lambda x: str(x[0])):
                censo_items.append({
                    "nombre": r_nom,
                    "nombre_display": simplificar_red_web(r_nom),
                    "cantidad": len(lista),
                    "tipo_label": "Centros"
                })
        elif red_param and not centro_param:
            # Nivel 2: Red específica -> Centros de Salud
            r_corta = simplificar_red_web(red_param)
            censo_titulo = f"🏥 Distribución de {tipo_label_censo} por Centro de Salud — {r_corta}"
            censo_subtitulo = f"Total en esta Red: {total:,} registros | Haz clic en un Centro para ver sus Áreas"
            grupos = {}
            for a in activos_contexto:
                c_nom = a.get('centro_salud_nombre') or 'Sin Centro'
                grupos.setdefault(c_nom, []).append(a)
            for c_nom, lista in sorted(grupos.items(), key=lambda x: len(x[1]), reverse=True):
                censo_items.append({
                    "nombre": c_nom,
                    "nombre_display": c_nom,
                    "cantidad": len(lista),
                    "tipo_label": "Áreas"
                })
        else:
            # Nivel 3: Centro de Salud específico -> Áreas
            censo_titulo = f"📍 Distribución de {tipo_label_censo} por Área / Servicio — {centro_param}"
            censo_subtitulo = f"Total en este Centro: {total:,} registros | Haz clic en un Área para listar sus activos"
            grupos = {}
            for a in activos_contexto:
                a_nom = a.get('area') or a.get('servicio') or 'General'
                grupos.setdefault(a_nom, []).append(a)
            for a_nom, lista in sorted(grupos.items(), key=lambda x: len(x[1]), reverse=True):
                censo_items.append({
                    "nombre": a_nom,
                    "nombre_display": a_nom,
                    "cantidad": len(lista),
                    "tipo_label": "Lista de Activos"
                })

        # Datos para gráficas Chart.js
        conteo_tipos = Counter([a.get("nombre", "Activo").strip() for a in activos_contexto if a.get("nombre")])
        top_tipos = conteo_tipos.most_common(8)
        chart_tipos_labels = json.dumps([k for k, v in top_tipos])
        chart_tipos_data = json.dumps([v for k, v in top_tipos])

        chart_censo_labels = json.dumps([item["nombre_display"] for item in censo_items[:8]])
        chart_censo_data = json.dumps([item["cantidad"] for item in censo_items[:8]])

        activos_json = json.dumps(activos_contexto, default=str)

        return render_template_string(
            HTML_ANALISIS,
            redes=redes_db,
            centros=centros_db,
            red_sel=red_param,
            centro_sel=centro_param,
            tipo_sel=tipo_param,
            cnt_todo=cnt_todo,
            cnt_equipos=cnt_equipos,
            cnt_muebles=cnt_muebles,
            total=total,
            operativos=operativos,
            garantia=garantia,
            bajas=bajas,
            censo_titulo=censo_titulo,
            censo_subtitulo=censo_subtitulo,
            censo_items=censo_items,
            chart_censo_labels=chart_censo_labels,
            chart_censo_data=chart_censo_data,
            chart_tipos_labels=chart_tipos_labels,
            chart_tipos_data=chart_tipos_data,
            eqs_vista=activos_contexto,
            activos_json=activos_json
        )
    except Exception as e:
        return f"Error en análisis: {e}", 500

# =========================================================================
# VISTA WEB PÚBLICA: ENLACES Y ACCESOS DEL SISTEMA
# =========================================================================
HTML_ENLACES_PORTAL = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>SGEM GAMLP • Enlaces y Accesos del Sistema</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #0F172A;
            --primary-light: #1E293B;
            --accent: #0284C7;
            --bg: #F8FAFC;
            --card: #FFFFFF;
            --text: #0F172A;
            --muted: #64748B;
            --border: #E2E8F0;
            --radius: 12px;
        }
        * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; }
        body { background-color: var(--bg); color: var(--text); padding-bottom: 50px; -webkit-font-smoothing: antialiased; }
        .header {
            background: #0F172A;
            color: #FFFFFF;
            padding: 24px 20px 20px;
            text-align: center;
            border-bottom: 1px solid #1E293B;
            box-shadow: 0 1px 3px rgba(0,0,0,0.1);
        }
        .badge-gamlp { display: inline-block; background: rgba(255,255,255,0.08); border: 1px solid rgba(255,255,255,0.14); padding: 3px 12px; border-radius: 20px; font-size: 11px; margin-bottom: 8px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.5px; }
        .header h1 { font-size: 22px; font-weight: 800; margin-bottom: 4px; letter-spacing: -0.5px; }
        .header p { font-size: 13px; color: #94A3B8; }
        .nav-tabs { display: flex; justify-content: center; gap: 8px; margin-top: 15px; flex-wrap: wrap; }
        .nav-tab {
            color: #94A3B8;
            text-decoration: none;
            padding: 6px 14px;
            border-radius: 20px;
            font-size: 12.5px;
            font-weight: 600;
            background: rgba(255,255,255,0.06);
            border: 1px solid rgba(255,255,255,0.1);
            transition: all 0.2s;
        }
        .nav-tab:hover { border-color: rgba(255,255,255,0.25); color: #FFFFFF; }
        .nav-tab.active { background: #FFFFFF; color: #0F172A; font-weight: 700; box-shadow: 0 2px 8px rgba(0,0,0,0.2); }
        .container { max-width: 860px; margin: 20px auto 0; padding: 0 16px; }
        .banner {
            background: #0F172A;
            border: 1px solid #1E293B;
            color: white;
            padding: 20px;
            border-radius: var(--radius);
            margin-bottom: 20px;
            box-shadow: 0 2px 8px rgba(15,23,42,0.1);
        }
        .link-card {
            background: var(--card);
            border-radius: var(--radius);
            border: 1px solid var(--border);
            padding: 18px 20px;
            margin-bottom: 14px;
            box-shadow: 0 2px 8px rgba(0,0,0,0.03);
            transition: transform 0.15s ease, box-shadow 0.15s ease;
        }
        .link-card:hover { transform: translateY(-2px); box-shadow: 0 6px 16px rgba(0,0,0,0.06); }
        .card-header { display: flex; align-items: center; justify-content: space-between; margin-bottom: 6px; }
        .card-title { font-size: 16px; font-weight: 800; color: var(--primary); display: flex; align-items: center; gap: 8px; }
        .badge { font-size: 11px; font-weight: 700; padding: 3px 8px; border-radius: 4px; }
        .card-desc { font-size: 13px; color: var(--muted); line-height: 1.45; margin-bottom: 10px; }
        .url-box {
            background: #F8FAFC;
            border: 1px solid #E2E8F0;
            border-radius: 8px;
            padding: 9px 12px;
            font-family: monospace;
            font-size: 13px;
            font-weight: 700;
            color: #0284C7;
            word-break: break-all;
            margin-bottom: 12px;
        }
        .btn-group { display: flex; gap: 8px; flex-wrap: wrap; }
        .btn {
            padding: 8px 14px;
            border-radius: 8px;
            font-size: 12.5px;
            font-weight: 700;
            border: none;
            cursor: pointer;
            text-decoration: none;
            display: inline-flex;
            align-items: center;
            gap: 5px;
            transition: all 0.2s;
        }
        .btn-copy { background: #2563EB; color: white; }
        .btn-wa { background: #22C55E; color: white; }
        .btn-open { background: #F1F5F9; color: #003B64; border: 1px solid #CBD5E1; }
    </style>
</head>
<body>
    <div class="header">
        <span class="badge-gamlp">GAMLP • SGEM v1.1</span>
        <h1>Sistema de Gestión de Equipamiento Médico</h1>
        <p>Directorio Central de Enlaces y Accesos</p>
        <div class="nav-tabs">
            <a href="/inventario" class="nav-tab">📦 Inventario</a>
            <a href="/analisis" class="nav-tab">📊 Análisis y Censo</a>
            <a href="/movil" class="nav-tab">📱 Registro Móvil</a>
            <a href="/historico" class="nav-tab">🏛️ Histórico</a>
            <a href="/enlaces" class="nav-tab active">🔗 Enlaces</a>
        </div>
    </div>

    <div class="container">
        <div class="banner">
            <h2 style="font-size: 20px; font-weight: 800; margin-bottom: 6px;">🔗 Directorio de Enlaces y Accesos Oficiales</h2>
            <p style="font-size: 13.5px; opacity: 0.95; line-height: 1.5;">
                Todos los accesos al sistema SGEM GAMLP disponibles 24/7 en la nube, red interna y consulta histórica. Copia con un clic o comparte directamente por WhatsApp.
            </p>
        </div>

        <!-- 1. SUITE MÓVIL (NUBE 24/7 - RELEVAMIENTO 2026) -->
        <div class="link-card" style="border-left: 5px solid #16A34A;">
            <div class="card-header">
                <div class="card-title">📱 Suite Móvil - Relevamiento 2026 (Actual)</div>
                <span class="badge" style="background: #DCFCE7; color: #166534;">🟢 ACTIVO ONLINE OFICIAL</span>
            </div>
            <div class="card-desc">
                Acceso para trabajo de campo en celulares y tablets. Registro de nuevos equipos médicos, muebles, áreas y mantenimientos en la base de datos oficial 2026.
            </div>
            <div class="url-box">https://cmms-gamlp.onrender.com/movil</div>
            <div class="btn-group">
                <button class="btn btn-copy" onclick="copiarTexto('https://cmms-gamlp.onrender.com/movil', this)">📋 Copiar Enlace</button>
                <a href="https://api.whatsapp.com/send?text=Acceso%20Suite%20M%C3%B3vil%20SGEM%20GAMLP%3A%20https%3A%2F%2Fcmms-gamlp.onrender.com%2Fmovil" target="_blank" class="btn btn-wa">📲 Enviar por WhatsApp</a>
                <a href="https://cmms-gamlp.onrender.com/movil" target="_blank" class="btn btn-open">🚀 Abrir Relevamiento</a>
            </div>
        </div>

        <!-- 2. CONSULTA HISTÓRICA (GESTIONES ANTERIORES - SOLO LECTURA) -->
        <div class="link-card" style="border-left: 5px solid #D97706;">
            <div class="card-header">
                <div class="card-title">🏛️ Consulta Histórica GAMLP (Gestiones Anteriores - Solo Lectura)</div>
                <span class="badge" style="background: #FEF3C7; color: #92400E;">🏛️ SOLO LECTURA • 2.938 EQUIPOS</span>
            </div>
            <div class="card-desc">
                Consulta y búsqueda de antecedentes de los 2.938 equipos médicos y mobiliario relevados en gestiones pasadas. Fichas técnicas completas protegidas contra modificación o borrado accidental.
            </div>
            <div class="url-box">https://cmms-gamlp.onrender.com/historico</div>
            <div class="btn-group">
                <button class="btn btn-copy" style="background: #D97706; color: white;" onclick="copiarTexto('https://cmms-gamlp.onrender.com/historico', this)">📋 Copiar Enlace</button>
                <a href="https://api.whatsapp.com/send?text=Consulta%20Hist%C3%B3rica%20SGEM%20GAMLP%3A%20https%3A%2F%2Fcmms-gamlp.onrender.com%2Fhistorico" target="_blank" class="btn btn-wa">📲 Enviar por WhatsApp</a>
                <a href="https://cmms-gamlp.onrender.com/historico" target="_blank" class="btn btn-open">🚀 Abrir Consulta Histórica</a>
            </div>
        </div>

        <!-- 3. PORTAL WEB GENERAL (NUBE 24/7) -->
        <div class="link-card" style="border-left: 5px solid #005691;">
            <div class="card-header">
                <div class="card-title">🌐 Portal Web General GAMLP (Nube 24/7)</div>
                <span class="badge" style="background: #E0F2FE; color: #075985;">🌐 PÁGINA PRINCIPAL</span>
            </div>
            <div class="card-desc">
                Portal institucional de bienvenida, módulos informativos, inventario descentralizado y análisis territorial de la red hospitalaria.
            </div>
            <div class="url-box">https://cmms-gamlp.onrender.com/</div>
            <div class="btn-group">
                <button class="btn btn-copy" onclick="copiarTexto('https://cmms-gamlp.onrender.com/', this)">📋 Copiar Enlace</button>
                <a href="https://cmms-gamlp.onrender.com/" target="_blank" class="btn btn-open">🚀 Abrir Portal</a>
            </div>
        </div>

        <!-- 4. RED LOCAL WI-FI -->
        <div class="link-card" style="border-left: 5px solid #64748B;">
            <div class="card-header">
                <div class="card-title">📶 Conexión Red Local Wi-Fi (Servidor Oficina)</div>
                <span class="badge" style="background: #F1F5F9; color: #475569;">🏢 RED LOCAL INTERNA</span>
            </div>
            <div class="card-desc">
                Para ingresar desde celulares o computadoras conectadas a la misma red Wi-Fi de la oficina mientras el software de escritorio esté abierto en la PC principal.
            </div>
            <div class="url-box" id="lbl_local">http://&lt;IP_DE_TU_PC&gt;:5000/movil</div>
            <div class="btn-group">
                <button class="btn" style="background: #475569; color: white;" onclick="copiarTexto(document.getElementById('lbl_local').innerText.trim(), this)">📋 Copiar Formato</button>
                <button class="btn btn-open" onclick="detectarIp()">🔍 Detectar Dirección Actual</button>
            </div>
        </div>

    </div>

    <script>
        function copiarTexto(texto, btn) {
            if (navigator.clipboard && window.isSecureContext) {
                navigator.clipboard.writeText(texto).then(() => feedback(btn)).catch(() => fallback(texto, btn));
            } else {
                fallback(texto, btn);
            }
        }
        function fallback(texto, btn) {
            const el = document.createElement('input');
            el.value = texto;
            document.body.appendChild(el);
            el.select();
            try {
                document.execCommand('copy');
                feedback(btn);
            } catch(e) {
                prompt('Copia el enlace manualmente:', texto);
            }
            document.body.removeChild(el);
        }
        function feedback(btn) {
            if (!btn) return;
            const orig = btn.innerHTML;
            btn.innerHTML = '✅ ¡Copiado!';
            const origBg = btn.style.backgroundColor;
            btn.style.backgroundColor = '#16A34A';
            setTimeout(() => {
                btn.innerHTML = orig;
                btn.style.backgroundColor = origBg;
            }, 2000);
        }
        function detectarIp() {
            const u = window.location.origin + '/movil';
            const l = document.getElementById('lbl_local');
            if (l) l.innerText = u;
            alert('Dirección detectada:\\n' + u);
        }
    </script>
</body>
</html>
"""

@app_web.route('/enlaces')
def vista_enlaces_web():
    """Directorio web de enlaces y accesos institucionales."""
    return render_template_string(HTML_ENLACES_PORTAL)

@app_web.route('/historico')
@app_web.route('/movil/historico')
def vista_historica_web():
    """Portal web de consulta histórica GAMLP (Solo Lectura • Gestiones Anteriores)."""
    return render_template_string(HTML_HISTORICO_WEB)

@app_web.route('/descargar/<filename>')
def descargar_archivo(filename):
    # Si es PDF, se sirve inline (para abrir directamente en navegador móvil sin bloquear descargas)
    attachment_flag = not filename.lower().endswith('.pdf')
    # Buscar el archivo de forma recursiva en las carpetas de las Áreas
    for root, dirs, files in os.walk(CARPETAS["areas"]):
        if filename in files:
            return send_from_directory(root, filename, as_attachment=attachment_flag)
    # Fallback
    return send_from_directory(CARPETAS["areas"], filename, as_attachment=attachment_flag)

@app_web.route('/equipo/<path:id_equipo>/descargar_qr')
def descargar_qr_web(id_equipo):
    import urllib.parse
    id_equipo = urllib.parse.unquote(str(id_equipo)).strip()
    try:
        conn = obtener_conexion()
        if not conn:
            return "Error de conexión", 500
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("SELECT * FROM equipos WHERE id = %s OR id = %s OR id = %s", (id_equipo, id_equipo.replace('S/N ', 'SN-'), id_equipo.replace('SN-', 'S/N ')))
        eq = cur.fetchone()
        cur.close()
        conn.close()
        
        if not eq:
            return "Equipo no encontrado", 404
            
        import qrcode
        from PIL import Image, ImageDraw, ImageFont
        import io
        from flask import send_file
        
        url_base = os.environ.get("RENDER_EXTERNAL_URL") or CONFIG.get("url_base_web", "https://cmms-gamlp.onrender.com")
        enl = f"{url_base}/equipo/{urllib.parse.quote(str(eq['id']))}"
        
        qr_base = qrcode.QRCode(version=1, box_size=12, border=1)
        qr_base.add_data(enl)
        qr_base.make(fit=True)
        img_qr_pil = qr_base.make_image(fill_color="black", back_color="white").convert("RGB")
        
        qr_w, qr_h = img_qr_pil.size
        extra_h = 100
        total_w = qr_w
        total_h = qr_h + extra_h
        
        sticker = Image.new("RGB", (total_w, total_h), "white")
        sticker.paste(img_qr_pil, (0, 0))
        
        draw = ImageDraw.Draw(sticker)
        
        try:
            font_nom = ImageFont.truetype("arial.ttf", 15)
            font_id = ImageFont.truetype("arialbd.ttf", 16)
            font_loc = ImageFont.truetype("arial.ttf", 14)
        except:
            font_nom = ImageFont.load_default()
            font_id = ImageFont.load_default()
            font_loc = ImageFont.load_default()
            
        txt_nombre = eq['nombre']
        txt_id = eq['id']
        txt_area = eq.get('area', 'General')
        
        def draw_centered_text(text, y_pos, font, color="black"):
            try:
                w = draw.textlength(text, font=font)
            except:
                try:
                    w = draw.textsize(text, font=font)[0]
                except:
                    w = len(text) * 8
            x_pos = (total_w - w) / 2
            draw.text((x_pos, y_pos), text, fill=color, font=font)
            
        draw_centered_text(txt_nombre, qr_h + 10, font_nom)
        draw_centered_text(txt_id, qr_h + 35, font_id)
        draw_centered_text(txt_area, qr_h + 60, font_loc)
        
        img_io = io.BytesIO()
        sticker.save(img_io, 'PNG')
        img_io.seek(0)
        
        id_sanitizado = "".join([c for c in id_equipo if c.isalnum() or c in ('-', '_')]).strip()
        return send_file(img_io, mimetype='image/png', as_attachment=True, download_name=f"QR_{id_sanitizado}.png")
    except Exception as e:
        return f"Error al generar QR: {e}", 500

HTML_FICHA_MUEBLE = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ mueble['descripcion'] or 'Ficha de Mueble y TI' }} - SGEM GAMLP</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #0F172A;
            --primary-dark: #020617;
            --primary-hover: #1E293B;
            --bg: #F8FAFC;
            --card-bg: #FFFFFF;
            --text: #0F172A;
            --muted: #64748B;
            --border: #E2E8F0;
            --border-hover: #CBD5E1;
            --green: #10B981;
            --orange: #F59E0B;
            --red: #EF4444;
        }

        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background-color: var(--bg);
            color: var(--text);
            margin: 0;
            padding: 0;
            -webkit-font-smoothing: antialiased;
        }

        .header {
            background: #0F172A;
            color: white;
            padding: 24px 20px 20px;
            text-align: center;
            border-bottom: 1px solid #1E293B;
            box-shadow: 0 1px 3px rgba(0,0,0,0.1);
        }

        .badge-gamlp {
            display: inline-block;
            background: rgba(255, 255, 255, 0.08);
            border: 1px solid rgba(255, 255, 255, 0.14);
            color: #CBD5E1;
            font-size: 11px;
            font-weight: 700;
            padding: 3px 12px;
            border-radius: 20px;
            margin-bottom: 8px;
            letter-spacing: 0.5px;
            text-transform: uppercase;
        }

        .header h1 {
            font-size: 20px;
            font-weight: 800;
            margin: 0 0 4px;
            letter-spacing: -0.5px;
        }

        .header p {
            font-size: 13px;
            color: #94A3B8;
            margin: 0;
        }

        .container {
            max-width: 900px;
            margin: 0 auto;
            padding: 24px 16px 60px;
        }

        .top-nav {
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 10px;
            margin-bottom: 20px;
        }

        .btn-nav {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            padding: 8px 16px;
            background: white;
            color: var(--text);
            border: 1.5px solid var(--border);
            border-radius: 10px;
            text-decoration: none;
            font-size: 13px;
            font-weight: 700;
            transition: all 0.2s;
            cursor: pointer;
        }

        .btn-nav:hover {
            border-color: var(--primary);
            color: var(--primary);
        }

        .btn-nav.primary {
            background: var(--primary);
            color: white;
            border-color: var(--primary);
        }

        .main-card {
            background: white;
            border: 1px solid var(--border);
            border-radius: 18px;
            padding: 24px;
            box-shadow: 0 4px 16px rgba(0,0,0,0.04);
            margin-bottom: 24px;
        }

        .title-row {
            display: flex;
            align-items: flex-start;
            gap: 16px;
            border-bottom: 1px solid var(--border);
            padding-bottom: 20px;
            margin-bottom: 20px;
        }

        .title-icon {
            font-size: 40px;
            background: #F1F5F9;
            width: 64px;
            height: 64px;
            display: flex;
            align-items: center;
            justify-content: center;
            border-radius: 14px;
            flex-shrink: 0;
        }

        .title-info h2 {
            font-size: 20px;
            font-weight: 800;
            color: var(--text);
            margin: 0 0 8px;
            line-height: 1.3;
        }

        .badges-row {
            display: flex;
            flex-wrap: wrap;
            gap: 8px;
            align-items: center;
        }

        .badge {
            display: inline-block;
            padding: 4px 10px;
            border-radius: 12px;
            font-size: 12px;
            font-weight: 700;
        }

        .badge-bueno { background: #ECFDF5; color: #065F46; border: 1px solid #A7F3D0; }
        .badge-regular { background: #FEF3C7; color: #92400E; border: 1px solid #FDE68A; }
        .badge-baja { background: #FEE2E2; color: #991B1B; border: 1px solid #FECACA; }
        .badge-tag { background: #EFF6FF; color: #1D4ED8; border: 1px solid #BFDBFE; }

        .info-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(360px, 1fr));
            gap: 20px;
        }

        .section-box {
            background: #F8FAFC;
            border: 1px solid var(--border);
            border-radius: 14px;
            padding: 16px 18px;
        }

        .section-box h3 {
            font-size: 14px;
            font-weight: 800;
            color: #334155;
            margin: 0 0 12px;
            display: flex;
            align-items: center;
            gap: 6px;
            border-bottom: 1px solid #E2E8F0;
            padding-bottom: 8px;
        }

        .data-table {
            width: 100%;
            border-collapse: collapse;
            font-size: 13px;
        }

        .data-table tr {
            border-bottom: 1px dashed #E2E8F0;
        }

        .data-table tr:last-child {
            border-bottom: none;
        }

        .data-table td {
            padding: 8px 4px;
            vertical-align: top;
        }

        .data-table td.label {
            font-weight: 700;
            color: #64748B;
            width: 42%;
        }

        .data-table td.value {
            font-weight: 600;
            color: var(--text);
        }

        .obs-box {
            background: #FFFBEB;
            border: 1px solid #FDE68A;
            border-radius: 12px;
            padding: 14px 16px;
            font-size: 13px;
            color: #78350F;
            line-height: 1.5;
            margin-top: 10px;
        }

        @media print {
            .no-print { display: none !important; }
            body { background: white; color: black; }
            .container { max-width: 100%; padding: 0; }
            .main-card { border: 1px solid #333; box-shadow: none; border-radius: 0; padding: 15px; }
            .section-box { border: 1px solid #888; background: #fafafa; border-radius: 0; }
        }
    </style>
</head>
<body>
    <div class="header no-print">
        <span class="badge-gamlp">GAMLP • SGEM v1.1</span>
        <h1>Ficha Técnica de Bien y Mueblería</h1>
        <p>Gobierno Autónomo Municipal de La Paz</p>
    </div>

    <div class="container">
        <div class="top-nav no-print">
            <div style="display: flex; gap: 8px;">
                <a href="/inventario" class="btn-nav">📦 Volver al Inventario</a>
                <a href="/analisis" class="btn-nav">📊 Censo y Análisis</a>
            </div>
            <button onclick="window.print()" class="btn-nav primary">🖨️ Imprimir / Guardar PDF</button>
        </div>

        <div class="main-card">
            <div class="title-row">
                <div class="title-icon">{{ icono }}</div>
                <div class="title-info" style="flex: 1;">
                    <h2>{{ mueble['descripcion'] or 'Activo / Bien Municipal' }}</h2>
                    <div class="badges-row">
                        {% set est = (mueble['estado_conservacion'] or mueble['estado'] or 'Bueno')|lower %}
                        {% if 'baja' in est or 'mal' in est %}
                        <span class="badge badge-baja">🛑 {{ mueble['estado_conservacion'] or 'Baja' }}</span>
                        {% elif 'reg' in est or 'man' in est %}
                        <span class="badge badge-regular">⚠️ {{ mueble['estado_conservacion'] or 'Regular' }}</span>
                        {% else %}
                        <span class="badge badge-bueno">✅ {{ mueble['estado_conservacion'] or 'Bueno' }}</span>
                        {% endif %}

                        <span class="badge badge-tag">🏷️ {{ mueble['tipo_activo'] or 'Mueble / TI' }}</span>
                        <span class="badge badge-tag">🆔 Cod. AF: {{ codigo_principal }}</span>
                    </div>
                </div>
            </div>

            <div class="info-grid">
                <!-- Tarjeta 1: Especificaciones Técnicas -->
                <div class="section-box">
                    <h3>⚙️ Datos Técnicos y Códigos</h3>
                    <table class="data-table">
                        <tr><td class="label">Código SISFAM:</td><td class="value"><strong>{{ mueble['codigo_sispam'] or 'S/C' }}</strong></td></tr>
                        <tr><td class="label">Código Bertin:</td><td class="value">{{ mueble['bertin'] or '-' }}</td></tr>
                        <tr><td class="label">Código SAPM:</td><td class="value">{{ mueble['sapm'] or '-' }}</td></tr>
                        <tr><td class="label">Tipo de Activo:</td><td class="value">{{ mueble['tipo_activo'] or 'Mueble / TI' }}</td></tr>
                        <tr><td class="label">Marca:</td><td class="value">{{ mueble['marca'] or 'Sin Marca' }}</td></tr>
                        <tr><td class="label">Modelo:</td><td class="value">{{ mueble['modelo'] or 'Sin Modelo' }}</td></tr>
                        <tr><td class="label">N° Serie:</td><td class="value">{{ mueble['serie'] or 'Sin Serie' }}</td></tr>
                        <tr><td class="label">Estado de Conservación:</td><td class="value">{{ mueble['estado_conservacion'] or 'Bueno' }}</td></tr>
                    </table>
                </div>

                <!-- Tarjeta 2: Asignación y Responsabilidad -->
                <div class="section-box">
                    <h3>👤 Asignación y Custodia</h3>
                    <table class="data-table">
                        <tr><td class="label">Persona Responsable:</td><td class="value"><strong>{{ mueble['persona_asignada'] or 'No asignado' }}</strong></td></tr>
                        <tr><td class="label">C.I. Asignado:</td><td class="value">{{ mueble['ci_asignado'] or '-' }}</td></tr>
                        <tr><td class="label">Cargo:</td><td class="value">{{ mueble['cargo_asignado'] or '-' }}</td></tr>
                        <tr><td class="label">Técnico Inventariador:</td><td class="value">{{ mueble['tecnico_inventareador'] or '-' }}</td></tr>
                        <tr><td class="label">Fecha Asignación:</td><td class="value">{{ mueble['fecha_asignacion'] or '-' }}</td></tr>
                        <tr><td class="label">Detalle Transacción:</td><td class="value">{{ mueble['detalle_transaccion'] or '-' }}</td></tr>
                    </table>
                </div>

                <!-- Tarjeta 3: Ubicación y Territorio -->
                <div class="section-box">
                    <h3>📍 Ubicación y Territorio</h3>
                    <table class="data-table">
                        <tr><td class="label">Red de Salud:</td><td class="value"><strong>{{ red_nombre }}</strong></td></tr>
                        <tr><td class="label">Centro de Salud:</td><td class="value"><strong>{{ centro_nombre }}</strong></td></tr>
                        <tr><td class="label">Unidad Organizacional:</td><td class="value">{{ mueble['unidad_organizacional'] or '-' }}</td></tr>
                        <tr><td class="label">Sector Actual:</td><td class="value">{{ mueble['sector_actual'] or '-' }}</td></tr>
                        <tr><td class="label">Ubicación / Área:</td><td class="value">{{ mueble['ubicacion'] or '-' }}</td></tr>
                    </table>
                </div>

                <!-- Tarjeta 4: Registro y Observaciones -->
                <div class="section-box">
                    <h3>📋 Registro y Notas</h3>
                    <table class="data-table">
                        <tr><td class="label">Fecha Incorporación:</td><td class="value">{{ mueble['fecha_incorporacion'] or '-' }}</td></tr>
                        <tr><td class="label">Fecha Registro:</td><td class="value">{{ mueble['fecha_registro'] or '-' }}</td></tr>
                    </table>
                    <div style="margin-top: 10px;">
                        <span style="font-size: 12px; font-weight: 700; color: #64748B;">Observaciones:</span>
                        <div class="obs-box">
                            {{ mueble['observaciones_de_asignacion'] or 'Sin observaciones registradas.' }}
                        </div>
                    </div>
                </div>
            </div>
        </div>
    </div>
</body>
</html>
"""

@app_web.route('/mueble/<path:mueble_id>')
def ver_mueble(mueble_id):
    clean_id = str(mueble_id).strip()
    if clean_id.upper().startswith("MUE-"):
        clean_id = clean_id[4:].strip()
    try:
        mid = int(clean_id)
    except ValueError:
        return "<h1>❌ ID de mueble inválido</h1>", 400

    try:
        conn = obtener_conexion()
        if not conn:
            return "<h1>❌ Error al conectar a la base de datos</h1>", 500
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("""
            SELECT m.*, r.nombre as red_nombre_fk, c.nombre as centro_nombre_fk
            FROM muebleria m
            LEFT JOIN redes_salud r ON m.red_salud_id = r.id
            LEFT JOIN centros_salud c ON m.centro_salud_id = c.id
            WHERE m.id = %s
        """, (mid,))
        m = cur.fetchone()
        
        cur.execute("SELECT id, nombre FROM redes_salud")
        redes_db = {r['id']: r['nombre'] for r in cur.fetchall()}
        
        cur.execute("SELECT id, nombre, red_salud_id FROM centros_salud")
        centros_db = [dict(r) for r in cur.fetchall()]
        cur.close()
        conn.close()

        if not m:
            return "<h1>❌ Registro de Mueble / TI no encontrado</h1>", 404

        mueble = dict(m)

        # Resolver centro y red si son nulos en la tabla
        cen_nom = mueble.get('centro_nombre_fk') or mueble.get('unidad_organizacional') or ''
        if not cen_nom:
            ub = mueble.get('ubicacion') or ''
            for c in centros_db:
                if c['nombre'].upper() in ub.upper():
                    cen_nom = c['nombre']
                    break
        if not cen_nom:
            cen_nom = 'Centro de Salud GAMLP'

        red_nom = mueble.get('red_nombre_fk') or mueble.get('direccion_administrativa') or ''
        if not red_nom:
            for c in centros_db:
                if c['nombre'].upper() == cen_nom.upper():
                    red_nom = redes_db.get(c['red_salud_id'], '')
                    break
        if not red_nom:
            red_nom = 'Red GAMLP'

        desc_lower = (str(mueble.get('descripcion') or '') + ' ' + str(mueble.get('tipo_activo') or '')).lower()
        if any(w in desc_lower for w in ['compu', 'monitor', 'laptop', 'pc', 'teclado', 'cpu', 'servidor', 'impresora']):
            icono = '💻'
        elif any(w in desc_lower for w in ['desfib', 'desfrib', 'electro', 'aspirad', 'monitor de signos', 'tensio', 'oxim']):
            icono = '🩺'
        else:
            icono = '🛋️'

        cod_af = mueble.get('codigo_sispam')
        if not cod_af or cod_af == 'S/C':
            cod_af = mueble.get('bertin') or mueble.get('sapm') or f"MUE-{mueble['id']:03d}"

        return render_template_string(
            HTML_FICHA_MUEBLE,
            mueble=mueble,
            centro_nombre=cen_nom,
            red_nombre=red_nom,
            icono=icono,
            codigo_principal=cod_af
        )
    except Exception as e:
        return f"<h1>Error al cargar ficha de mueble: {e}</h1>", 500

@app_web.route('/equipo/<path:id_equipo>')
def ver_equipo(id_equipo):
    import urllib.parse
    id_equipo = urllib.parse.unquote(str(id_equipo)).strip()
    try:
        conn = obtener_conexion()
        if not conn:
            return "<h1>❌ Error de conexión a Base de Datos</h1>", 500
            
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("SELECT * FROM equipos WHERE id = %s OR id = %s OR id = %s", (id_equipo, id_equipo.replace('S/N ', 'SN-'), id_equipo.replace('SN-', 'S/N ')))
        eq = cur.fetchone()
        
        if not eq:
            cur.close()
            conn.close()
            return "<h1>❌ Equipo no encontrado</h1>", 404
            
        cur.execute("SELECT * FROM historial_intervenciones WHERE equipo_id = %s OR equipo_id = %s ORDER BY fecha DESC", (eq['id'], id_equipo))
        historial = cur.fetchall()

        cat_str = f"{eq['nombre']} - {eq.get('marca') or ''} - {eq.get('modelo') or ''}"
        eq_nom = eq.get('nombre') or ''
        cur.execute("""
            SELECT nombre_repuesto, modelo_parte, cantidad, costo, estado_disponibilidad, observaciones
            FROM repuestos 
            WHERE (tipo_equipo = %s OR tipo_equipo = %s OR tipo_equipo ILIKE %s OR tipo_equipo ILIKE '%%general%%' OR tipo_equipo ILIKE '%%vacio%%')
            ORDER BY nombre_repuesto ASC
        """, (cat_str, eq_nom, f"%{eq_nom}%"))
        todos_rep = [dict(r) for r in cur.fetchall()]
        cur.close()
        conn.close()

        rep_stock = [r for r in todos_rep if str(r.get("estado_disponibilidad") or "En Stock").strip().lower() != "requerido" and int(r.get("cantidad") or 0) > 0]
        rep_req = [r for r in todos_rep if str(r.get("estado_disponibilidad") or "En Stock").strip().lower() == "requerido"]

        # Mapear archivos locales de hojas de trabajo existentes
        historial_list = []
        area_name = eq.get("area", "General")
        area_folder = "".join([c for c in area_name if c.isalnum() or c==' ']).strip()
        dir_mantenimiento = os.path.join(CARPETAS["areas"], area_folder, "mantenimientos")
        
        archivos_locales = []
        if os.path.exists(dir_mantenimiento):
            try:
                archivos_locales = os.listdir(dir_mantenimiento)
            except:
                pass
                
        for row in historial:
            d = dict(row)
            fecha_str = str(d['fecha'])
            fecha_compacta = fecha_str.replace('-', '')
            
            xlsx_match = None
            pdf_match = None
            id_sanitizado = sanitizar_nombre_archivo(id_equipo)
            prefix_1 = f"HT_{id_sanitizado}_{fecha_compacta}"
            prefix_alt = f"HT_{id_equipo}_{fecha_compacta}"
            
            for filename in archivos_locales:
                if filename.startswith(prefix_1) or filename.startswith(prefix_alt):
                    if filename.endswith(".xlsx"):
                        xlsx_match = filename
                    elif filename.endswith(".pdf"):
                        pdf_match = filename
            
            d['xlsx_file'] = xlsx_match
            d['pdf_file'] = pdf_match
            historial_list.append(d)

        # Calcular tiempo de garantía restante
        hoy = date.today()
        garantia_str = "Sin Garantía"
        if eq.get('garantia') == "Con Garantía" and eq.get('fecha_vencimiento_garantia'):
            f_venc = eq['fecha_vencimiento_garantia']
            if isinstance(f_venc, str):
                try:
                    f_venc = datetime.strptime(f_venc, "%Y-%m-%d").date()
                except:
                    f_venc = None
            if f_venc:
                if f_venc < hoy:
                    garantia_str = f"Vencida (Venció el {f_venc})"
                else:
                    from dateutil.relativedelta import relativedelta
                    diff = relativedelta(f_venc, hoy)
                    parts = []
                    if diff.years > 0:
                        parts.append(f"{diff.years} {'año' if diff.years == 1 else 'años'}")
                    if diff.months > 0:
                        parts.append(f"{diff.months} {'mes' if diff.months == 1 else 'meses'}")
                    if diff.days > 0:
                        parts.append(f"{diff.days} {'día' if diff.days == 1 else 'días'}")
                    duracion = ", ".join(parts) if parts else "Vence hoy"
                    garantia_str = f"Activa (Vence el {f_venc} - Resta: {duracion})"

        html_web = """
        <!DOCTYPE html><html lang="es"><head><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Ficha Técnica | {{ eq['nombre'] }}</title>
        <link rel="preconnect" href="https://fonts.googleapis.com">
        <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
        <style>
            * { box-sizing: border-box; }
            body { font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #F8FAFC; margin: 0; padding: 20px 15px; color: #0F172A; -webkit-font-smoothing: antialiased; }
            .tarjeta { background: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 14px; padding: 24px; box-shadow: 0 1px 3px rgba(0,0,0,0.02); max-width: 650px; margin: auto; }
            .cabecera { background: #0F172A; color: white; padding: 14px 18px; border-radius: 10px; text-align: center; margin-bottom: 20px; }
            .cabecera h3 { margin: 0; font-size: 14px; font-weight: 700; letter-spacing: 0.3px; }
            .estado { display: inline-block; padding: 4px 10px; border-radius: 6px; font-weight: 700; font-size: 12px; background: #ECFDF5; color: #065F46; border: 1px solid #A7F3D0; }
            .estado.baja { background: #FEE2E2; color: #991B1B; border: 1px solid #FECACA; }
            .estado.espera { background: #FEF3C7; color: #92400E; border: 1px solid #FDE68A; }
            .estado.inoperante { background: #F1F5F9; color: #475569; border: 1px solid #E2E8F0; }
            .btn-action { display: block; text-align: center; background: #0F172A; color: white; text-decoration: none; padding: 12px; border-radius: 8px; font-weight: 700; font-size: 13.5px; transition: all 0.2s; }
            .btn-action:hover { background: #1E293B; }
            .btn-action-sec { display: block; text-align: center; background: #FFFFFF; color: #0F172A; border: 1.5px solid #E2E8F0; text-decoration: none; padding: 11px; border-radius: 8px; font-weight: 700; font-size: 13.5px; transition: all 0.2s; }
            .btn-action-sec:hover { border-color: #0F172A; }
            table { width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 13px; }
            th, td { border-bottom: 1px solid #E2E8F0; padding: 10px 12px; text-align: left; }
            th { color: #64748B; font-size: 11px; font-weight: 700; text-transform: uppercase; background: #F8FAFC; }
        </style></head><body>
            <div id="sync-banner" style="display:none; background: #F59E0B; color: white; padding: 12px; text-align: center; font-weight: bold; font-size: 14px; border-radius: 8px; margin: 10px auto; max-width: 600px;">
                ⚠️ Tienes <span id="sync-count">0</span> reporte(s) guardado(s) offline. 
                <a href="#" onclick="intentarSincronizarAhora(); return false;" style="color: white; text-decoration: underline; margin-left: 10px;">Sincronizar ahora</a>
            </div>
            <div class="tarjeta"><div class="cabecera"><h3>SGEM GAMLP • Sistema de Gestión de Equipamiento Médico</h3></div>
                <h2 style="margin:0; font-size: 20px; font-weight: 800;">{{ eq['nombre'] }}</h2>
                <div style="margin-top: 6px;">
                    <div class="estado {% if eq['estado'] == 'Baja' %}baja{% elif eq['estado'] == 'En Espera de Repuesto' %}espera{% elif eq['estado'] == 'Fuera de Servicio' %}inoperante{% endif %}">{{ eq['estado'] }}</div>
                </div>
                
                {% if eq['foto'] %}
                <div style="text-align: center; margin: 15px 0;">
                    <img src="{{ eq['foto'] }}" alt="{{ eq['nombre'] }}" style="max-width: 100%; max-height: 250px; border-radius: 10px; object-fit: contain; box-shadow: 0 1px 3px rgba(0,0,0,0.04); border: 1px solid #E2E8F0;">
                </div>
                {% endif %}

                <div style="margin-top: 20px; font-size: 14px; line-height: 1.7; color: #334155;">
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">ID:</strong> {{ eq['id'] }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">Red de Salud:</strong> {{ eq['red_salud_nombre'] or 'GAMLP' }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">Centro de Salud:</strong> {{ eq['centro_salud_nombre'] or 'Centro de Salud' }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">S/N:</strong> {{ eq['numero_serie'] or 'Sin Serie' }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">Modelo:</strong> {{ eq['marca'] }} / {{ eq['modelo'] }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">Área / Servicio:</strong> {{ eq['servicio'] }} - {{ eq['area'] }}</p>
                    <p style="margin: 4px 0;"><strong style="color: #0F172A;">Garantía:</strong> {{ garantia_str }}</p>
                    <p style="margin: 4px 0; color: #D97706;"><strong style="color: #0F172A;">Criticidad:</strong> {{ eq['criticidad'] }}</p>
                </div>
                
                <div style="display: flex; flex-direction: column; gap: 8px; margin-top: 20px;">
                    <a href="/equipo/{{ eq['id'] }}/mantenimiento" class="btn-action" style="margin:0;">🛠️ Registrar Mantenimiento</a>
                    <a href="/equipo/{{ eq['id'] }}/descargar_ficha" class="btn-action-sec" style="margin:0;">📥 Descargar Ficha Técnica (.xlsx)</a>
                    <a href="/equipo/{{ eq['id'] }}/descargar_qr" class="btn-action-sec" style="margin:0;">📥 Descargar Código QR (Etiqueta)</a>
                </div>
                
                <h3 style="margin-top:25px; border-bottom: 1px solid #E2E8F0; padding-bottom: 6px; font-size: 15px; font-weight: 700;">Historial de Mantenimientos</h3>
                <table><tr><th>Fecha</th><th>Tipo</th><th>Realizado Por</th><th>Trabajo Realizado</th><th>Ficha</th></tr>
                {% for m in hist %}<tr>
                    <td>{{ m['fecha'] }}</td>
                    <td><strong>{{ m['tipo'] }}</strong></td>
                    <td>{{ m['realizado_por'] or 'Técnico' }}</td>
                    <td>{{ m['trabajo'] or m['detalle'] or 'Sin detalle' }}</td>
                    <td>
                        <a href="/equipo/{{ eq['id'] }}/mantenimiento/{{ m['id'] }}/descargar_excel" style="text-decoration:none; color:#1D4ED8; font-weight:bold; padding: 4px 10px; border-radius: 6px; background: #EFF6FF; border: 1px solid #BFDBFE; display: inline-block; font-size: 12px;" title="Descargar Hoja de Trabajo Excel">📥 Ficha (.xlsx)</a>
                    </td>
                </tr>
                {% else %}<tr><td colspan="5" style="text-align:center; color:#8E8E93;">Sin intervenciones registradas</td></tr>{% endfor %}
                </table>

                <h3 style="margin-top:25px; border-bottom: 2px solid #F2F2F7; padding-bottom: 5px; color: #34C759;">📦 Repuestos en Stock (Disponibles)</h3>
                <table><tr><th>Repuesto</th><th>Modelo / P/N</th><th>Stock</th><th>Costo (Bs.)</th></tr>
                {% for r in rep_stock %}<tr>
                    <td><strong>{{ r['nombre_repuesto'] }}</strong></td>
                    <td>{{ r['modelo_parte'] or '-' }}</td>
                    <td style="color:#34C759; font-weight:bold;">{{ r['cantidad'] }}</td>
                    <td>{{ "%.2f"|format(r['costo']|float) if r['costo'] else '-' }}</td>
                </tr>
                {% else %}<tr><td colspan="4" style="text-align:center; color:#8E8E93;">Sin repuestos en stock para este equipo</td></tr>{% endfor %}
                </table>

                <h3 style="margin-top:25px; border-bottom: 2px solid #F2F2F7; padding-bottom: 5px; color: #FF9500;">⚠️ Repuestos Requeridos (Necesarios)</h3>
                <table><tr><th>Repuesto Requerido</th><th>Modelo / P/N</th><th>Cant.</th><th>Motivo / Obs.</th></tr>
                {% for r in rep_req %}<tr>
                    <td><strong style="color:#FF9500;">{{ r['nombre_repuesto'] }}</strong></td>
                    <td>{{ r['modelo_parte'] or '-' }}</td>
                    <td>{{ r['cantidad'] }}</td>
                    <td>{{ r['observaciones'] or '-' }}</td>
                </tr>
                {% else %}<tr><td colspan="4" style="text-align:center; color:#8E8E93;">Sin requerimientos pendientes</td></tr>{% endfor %}
                </table></div>
            <script>
                function intentarSincronizarAhora() {
                    var pendientes = JSON.parse(localStorage.getItem("mantenimientos_pendientes") || "[]");
                    if (pendientes.length === 0) return;
                    
                    var banner = document.getElementById("sync-banner");
                    if (banner) banner.innerHTML = "🔄 Sincronizando reportes offline... (" + pendientes.length + " restantes)";
                    
                    var item = pendientes[0];
                    var passVal = item.web_pass;
                    if (!passVal) {
                        passVal = prompt("Introduce tu contraseña (" + (item.web_user || "usuario") + ") para sincronizar el reporte de " + (item.nombre_equipo || "mantenimiento") + ":");
                        if (!passVal) {
                            if (banner) banner.innerHTML = "⚠️ Sincronización pausada. Se requiere contraseña.";
                            return;
                        }
                    }

                    var params = new URLSearchParams();
                    for (var key in item) {
                        if (key !== "web_pass") {
                            params.append(key, item[key]);
                        }
                    }
                    params.append("web_pass", passVal);
                    
                    fetch(item.url_sincronizacion, {
                        method: "POST",
                        body: params
                    })
                    .then(response => {
                        if (response.ok) {
                            pendientes.shift();
                            localStorage.setItem("mantenimientos_pendientes", JSON.stringify(pendientes));
                            if (pendientes.length > 0) {
                                intentarSincronizarAhora();
                            } else {
                                if (banner) {
                                    banner.style.background = "#34C759";
                                    banner.innerHTML = "✅ ¡Todos los reportes sincronizados correctamente!";
                                }
                                setTimeout(() => {
                                    window.location.reload();
                                }, 1500);
                            }
                        } else {
                            if (banner) banner.innerHTML = "⚠️ Error de credenciales al sincronizar. Verifica tu usuario y contraseña.";
                        }
                    })
                    .catch(err => {
                        if (banner) banner.innerHTML = "⚠️ Servidor no responde. Sincronización pendiente (se enviará al conectar).";
                    });
                }


                window.addEventListener("load", function() {
                    var pendientes = JSON.parse(localStorage.getItem("mantenimientos_pendientes") || "[]");
                    if (pendientes.length > 0) {
                        var banner = document.getElementById("sync-banner");
                        var count = document.getElementById("sync-count");
                        if (banner && count) {
                            count.innerText = pendientes.length;
                            banner.style.display = "block";
                        }
                        intentarSincronizarAhora();
                    }
                });
            </script>
        </body></html>
        """
        return render_template_string(html_web, eq=eq, hist=historial_list, garantia_str=garantia_str, rep_stock=rep_stock, rep_req=rep_req)
    except Exception as e:
        return f"Error en el servidor web: {e}"

@app_web.route('/equipo/<path:id_equipo>/descargar_ficha')
@app_web.route('/equipo/<path:id_equipo>/ficha_excel')
def descargar_ficha_tecnica_excel(id_equipo):
    import urllib.parse
    id_equipo = urllib.parse.unquote(str(id_equipo)).strip()
    try:
        conn = obtener_conexion()
        if not conn:
            return "Error de conexión a base de datos", 500
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("SELECT * FROM equipos WHERE id = %s OR id = %s OR id = %s", (id_equipo, id_equipo.replace('S/N ', 'SN-'), id_equipo.replace('SN-', 'S/N ')))
        eq_act = cur.fetchone()
        
        cur.execute("SELECT * FROM repuestos")
        repuestos = cur.fetchall()
        cur.close()
        conn.close()

        if not eq_act:
            return "Equipo no encontrado", 404

        ruta_plantilla = obtener_ruta_plantilla("plantilla_ficha.xlsx")
        if not os.path.exists(ruta_plantilla):
            return "Plantilla no encontrada", 500

        wb = openpyxl.load_workbook(ruta_plantilla)
        hoja = wb.active

        def escribir(celda, valor):
            try:
                if valor is not None:
                    hoja[celda].value = valor
            except:
                pass

        def escribir_rcm(celda, texto):
            try:
                hoja[celda].value = texto if texto else ''
            except:
                pass

        # 1. Datos Generales de Identificación y Territorio
        escribir('K4', eq_act.get('red_salud_nombre', ''))
        escribir('K5', eq_act.get('centro_salud_nombre', ''))
        escribir('K8', eq_act.get('nombre', ''))
        escribir('H11', eq_act.get('area', ''))
        escribir('H12', eq_act.get('servicio', ''))
        escribir('H13', eq_act.get('marca', ''))
        escribir('H14', eq_act.get('modelo', ''))
        escribir('H15', str(eq_act.get('id', '')))
        escribir('H16', eq_act.get('procedencia', ''))
        escribir('H17', eq_act.get('fabricante', ''))
        escribir('H18', eq_act.get('garantia', ''))
        escribir('H19', eq_act.get('proveedor', ''))
        escribir('H20', eq_act.get('numero_serie', ''))
        escribir('H21', str(eq_act.get('anio_fab') or ''))
        escribir('H22', str(eq_act.get('fecha_adquisicion') or ''))

        # 2. Datos Técnicos Oficiales del Equipo
        escribir('E24', eq_act.get('voltaje', '') or '')
        escribir('G25', eq_act.get('corriente', '') or '')
        escribir('I26', eq_act.get('potencia', '') or '')
        escribir('H27', eq_act.get('vida_util', '') or eq_act.get('temperatura', '') or '')
        escribir('D28', eq_act.get('peso', '') or '')
        escribir('F29', eq_act.get('dimensiones', '') or '')
        escribir('H30', eq_act.get('bateria_respaldo', '') or eq_act.get('resolucion', '') or '')
        escribir('H31', eq_act.get('version_software', '') or eq_act.get('humedad', '') or '')
        escribir('H32', eq_act.get('suministro_gases', '') or '')

        # 3. Repuestos
        cat_str = f"{eq_act['nombre']} - {eq_act.get('marca', '')} - {eq_act.get('modelo', '')}"
        repuestos_equipo = [r for r in repuestos if r.get("tipo_equipo") == cat_str]
        for idx, r in enumerate(repuestos_equipo[:5]):
            cell_row = 34 + idx
            nom_rep = r.get("nombre_repuesto", "")
            cant_rep = r.get("cantidad", 0)
            txt_rep = f"{nom_rep}  (Cantidad: {cant_rep})" if cant_rep is not None else nom_rep
            escribir(f'C{cell_row}', txt_rep)

        # 4. Tecnología Predominante (X)
        escribir('S25', eq_act.get('t_elec', ''))
        escribir('S27', eq_act.get('t_elco', ''))
        escribir('S29', eq_act.get('t_mec', ''))
        escribir('Z25', eq_act.get('t_hid', ''))
        escribir('Z27', eq_act.get('t_neu', ''))
        escribir('Z29', eq_act.get('t_vap', ''))

        # 5. Tipo Adquisición y Tipo de Equipo (X)
        escribir('S33', eq_act.get('a_comp', ''))
        escribir('S35', eq_act.get('a_como', ''))
        escribir('S37', eq_act.get('a_don', ''))
        escribir('Y33', eq_act.get('te_fijo', ''))
        escribir('Y35', eq_act.get('te_mov', ''))
        escribir('Y37', eq_act.get('te_por', ''))

        # 6. Categorización
        cat_data = eq_act.get("categorizacion_detalle") or []
        if isinstance(cat_data, str):
            try: cat_data = json.loads(cat_data)
            except: cat_data = []
            
        for i in range(13):
            valor = str(cat_data[i]) if i < len(cat_data) else ""
            if valor in ("1", "I"): 
                escribir(f'AK{24+i}', 'X')
            elif valor in ("2", "II"): 
                escribir(f'AM{24+i}', 'X')
            elif valor in ("3", "III"): 
                escribir(f'AO{24+i}', 'X')

        try:
            puntajes_int = []
            for x in cat_data:
                if str(x).isdigit():
                    puntajes_int.append(int(x))
                elif str(x) == "I":
                    puntajes_int.append(1)
                elif str(x) == "II":
                    puntajes_int.append(2)
                elif str(x) == "III":
                    puntajes_int.append(3)
                else:
                    puntajes_int.append(0)
            puntaje_total = sum(puntajes_int)
        except:
            puntaje_total = 0

        if puntaje_total >= 30 or eq_act.get("criticidad") == "Riesgo Alto":
            escribir('AO37', 'X')
            escribir('AB38', "3 veces al año")
        elif puntaje_total >= 20 or eq_act.get("criticidad") == "Riesgo Medio":
            escribir('AM37', 'X')
            escribir('AB38', "2 veces al año")
        else:
            escribir('AK37', 'X')
            escribir('AB38', "1 vez al año")
            pass

        # 7. Tablas RCM y Observaciones
        escribir_rcm('B41', eq_act.get('contexto_operacional'))
        escribir_rcm('L41', eq_act.get('funciones_equipo'))
        escribir_rcm('V41', eq_act.get('acciones_preventivas'))
        escribir_rcm('AE41', eq_act.get('acciones_falla'))

        escribir_rcm('B49', eq_act.get('fallas_funcionales'))
        escribir_rcm('L49', eq_act.get('causas_fallo'))
        escribir_rcm('V49', eq_act.get('efectos_fallo'))
        escribir_rcm('AE49', eq_act.get('efecto_entorno'))

        escribir_rcm('B58', eq_act.get('observaciones'))

        # Insertar Foto si existe
        foto_path = eq_act.get('foto')
        if foto_path:
            try:
                foto_str = str(foto_path).strip()
                if foto_str.startswith("data:image") or len(foto_str) > 200:
                    foto_b64 = foto_str.split(",", 1)[1] if "," in foto_str else foto_str
                    img_bytes = base64.b64decode(foto_b64)
                    img_stream = io.BytesIO(img_bytes)
                    from openpyxl.drawing.image import Image as ExcelImage
                    img_excel = ExcelImage(img_stream)
                elif os.path.exists(foto_str):
                    from openpyxl.drawing.image import Image as ExcelImage
                    img_excel = ExcelImage(foto_str)
                else:
                    img_excel = None

                if img_excel:
                    img_excel.width = 220
                    img_excel.height = 220
                    hoja.add_image(img_excel, 'AA11')
            except Exception as ex_foto:
                print(f"[WARN] No se pudo incrustar la foto en Ficha: {ex_foto}")

        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)

        id_sanitizado = str(id_equipo).replace("/", "_").replace("\\", "_")
        return send_file(
            out_io,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"Ficha_Tecnica_{id_sanitizado}.xlsx"
        )
    except Exception as e:
        return f"Error generando Ficha Técnica Excel: {e}", 500

@app_web.route('/equipo/<path:id_equipo>/mantenimiento/<int:m_id>/descargar_excel')
@app_web.route('/mantenimiento/<int:m_id>/descargar_excel')
def descargar_hoja_trabajo_excel(m_id, id_equipo=None):
    try:
        conn = obtener_conexion()
        if not conn:
            return "Error de conexión", 500
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("SELECT * FROM historial_intervenciones WHERE id = %s", (m_id,))
        m = cur.fetchone()
        if not m:
            cur.close()
            conn.close()
            return "Mantenimiento no encontrado", 404
            
        cur.execute("SELECT * FROM equipos WHERE id = %s", (m['equipo_id'],))
        eq_data = cur.fetchone()
        cur.close()
        conn.close()
        
        if not eq_data:
            return "Equipo no encontrado", 404
            
        ruta_plantilla_ht = obtener_ruta_plantilla("plantilla_trabajo.xlsx")
        wb = openpyxl.load_workbook(ruta_plantilla_ht)
        ws = wb.active
        
        # Escribir campos de cabecera y equipo
        escribir_en_celda_segura(ws, 'P5', eq_data.get('red_salud_nombre', ''))
        escribir_en_celda_segura(ws, 'P6', eq_data.get('centro_salud_nombre', ''))
        escribir_en_celda_segura(ws, 'F11', eq_data.get('area', ''))
        escribir_en_celda_segura(ws, 'AA11', eq_data.get('servicio', ''))
        escribir_en_celda_segura(ws, 'J15', eq_data.get('nombre', ''))
        escribir_en_celda_segura(ws, 'AE15', str(eq_data.get('id', '')))
        escribir_en_celda_segura(ws, 'E17', eq_data.get('procedencia', ''))
        escribir_en_celda_segura(ws, 'AB17', str(eq_data.get('anio_fab', '')))
        escribir_en_celda_segura(ws, 'E19', eq_data.get('marca', ''))
        escribir_en_celda_segura(ws, 'AB19', eq_data.get('fabricante', ''))
        escribir_en_celda_segura(ws, 'F21', eq_data.get('modelo', ''))
        escribir_en_celda_segura(ws, 'AB21', eq_data.get('numero_serie', ''))
        
        # Fechas
        f_rec_raw = str(m.get('fecha') or date.today())
        f_ent_raw = str(m.get('fecha_entrega') or m.get('fecha') or date.today())
        h_ejec = str(m.get('hora_entrega') or datetime.now().strftime('%H:%M'))
        try:
            f_rec_dt = datetime.strptime(f_rec_raw, '%Y-%m-%d').date()
            f_rec_str = f_rec_dt.strftime('%d / %m / %Y')
        except:
            f_rec_str = f_rec_raw
            
        try:
            f_ent_dt = datetime.strptime(f_ent_raw, '%Y-%m-%d').date()
            f_ent_str = f_ent_dt.strftime('%d / %m / %Y')
        except:
            f_ent_str = f_ent_raw
            
        escribir_en_celda_segura(ws, 'M23', f_rec_str)
        escribir_en_celda_segura(ws, 'I62', f"{f_ent_str}  {h_ejec}")
        
        # Nombre del técnico firmante responsable
        escribir_en_celda_segura(ws, 'J64', m.get('realizado_por') or 'Técnico GAMLP')
        
        # Condición
        cond = m.get('condicion')
        if cond == "Óptimo": marcar_x(ws, 'P26')
        elif cond == "Aceptable": marcar_x(ws, 'W26')
        elif cond == "Crítica": marcar_x(ws, 'AC26')
        elif cond == "Inoperante": marcar_x(ws, 'AJ26')
        elif cond == "F/Servicio": marcar_x(ws, 'AP26')

        # Estado Físico
        est = m.get('estado_equipo')
        if est == "Óptimo": marcar_x(ws, 'O29')
        elif est == "Bueno": marcar_x(ws, 'U29')
        elif est == "Regular": marcar_x(ws, 'AB29')
        elif est == "Malo": marcar_x(ws, 'AH29')
        elif est == "Obsoleto": marcar_x(ws, 'AO29')

        # Tipo Mantenimiento
        tipo = m.get('tipo')
        if tipo == "Preventivo": marcar_x(ws, 'Q43')
        else: marcar_x(ws, 'AL43')

        # Textos largos
        escribir_texto_largo(ws, 'B33', m.get('deficiencia', ''))
        escribir_texto_largo(ws, 'B47', m.get('trabajo', ''))
        escribir_texto_largo(ws, 'B53', m.get('observaciones', ''))
        
        # Guardar en memoria y enviar al cliente
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        
        id_sanitizado = "".join([c for c in str(eq_data['id']) if c.isalnum() or c in ('-', '_')])
        return send_file(
            out_io,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"HT_{id_sanitizado}_OT{m_id}.xlsx"
        )
    except Exception as e:
        return f"Error generando Hoja de Trabajo Excel: {e}", 500

@app_web.route('/repuestos/descargar_excel')
def descargar_repuestos_excel_web():
    tipo = request.args.get('tipo', 'Stock')
    try:
        conn = obtener_conexion()
        if not conn:
            return "Error de conexión a base de datos", 500
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        
        tipo_filtro = str(tipo).strip().lower()
        if tipo_filtro in ["stock", "inventario", "en stock"]:
            cur.execute("SELECT * FROM repuestos WHERE estado_disponibilidad != 'Requerido' ORDER BY nombre_repuesto ASC")
            rep_list = [dict(r) for r in cur.fetchall()]
            nombre_descarga = f"Inventario_Repuestos_{datetime.now().strftime('%Y%m%d')}.xlsx"
        else:
            cur.execute("SELECT * FROM repuestos WHERE estado_disponibilidad = 'Requerido' ORDER BY nombre_repuesto ASC")
            rep_list = [dict(r) for r in cur.fetchall()]
            nombre_descarga = f"Repuestos_Requeridos_{datetime.now().strftime('%Y%m%d')}.xlsx"
            
        cur.close()
        conn.close()

        from generador_repuestos_excel import generar_excel_repuestos_wb
        wb, temp_files = generar_excel_repuestos_wb(rep_list, tipo=tipo)
        
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        
        for tf in temp_files:
            try: os.remove(tf)
            except: pass
            
        return send_file(
            out_io,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=nombre_descarga
        )
    except Exception as e:
        return f"Error generando Excel de Repuestos: {e}", 500

@app_web.route('/equipo/<path:id_equipo>/mantenimiento', methods=['GET', 'POST'])
def registrar_mantenimiento(id_equipo):
    import urllib.parse
    id_equipo = urllib.parse.unquote(str(id_equipo)).strip()
    error_msg = None
    try:
        conn = obtener_conexion()
        if not conn:
            return "<h1>❌ Error de conexión a Base de Datos</h1>", 500
            
        cur = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        cur.execute("SELECT * FROM equipos WHERE id = %s OR id = %s OR id = %s", (id_equipo, id_equipo.replace('S/N ', 'SN-'), id_equipo.replace('SN-', 'S/N ')))
        eq = cur.fetchone()
        
        if not eq:
            cur.close()
            conn.close()
            return "<h1>❌ Equipo no encontrado</h1>", 404

        # Cargar los repuestos disponibles con stock mayor a cero y compatibles con el tipo de equipo
        cat_str = f"{eq['nombre']} - {eq.get('marca') or ''} - {eq.get('modelo') or ''}"
        eq_nom = eq.get('nombre') or ''
        cur.execute("""
            SELECT nombre_repuesto, cantidad, tipo_equipo 
            FROM repuestos 
            WHERE cantidad > 0 
              AND COALESCE(estado_disponibilidad, 'En Stock') = 'En Stock'
              AND (tipo_equipo = %s OR tipo_equipo = %s OR tipo_equipo ILIKE %s)
            ORDER BY nombre_repuesto ASC
        """, (cat_str, eq_nom, f"%{eq_nom}%"))
        repuestos_list = [dict(r) for r in cur.fetchall()]
        if not repuestos_list:
            cur.execute("""
                SELECT nombre_repuesto, cantidad, tipo_equipo 
                FROM repuestos 
                WHERE cantidad > 0 AND COALESCE(estado_disponibilidad, 'En Stock') = 'En Stock'
                ORDER BY nombre_repuesto ASC
            """)
            repuestos_list = [dict(r) for r in cur.fetchall()]


        if request.method == 'POST':
            web_user = request.form.get('web_user', '').strip()
            web_pass = request.form.get('web_pass', '').strip()
            
            # Validar credenciales
            usuario_valido = login(web_user, web_pass)
            if not usuario_valido:
                error_msg = "Usuario o contraseña incorrectos."
            else:
                tipo = request.form.get('tipo')
                tipo_ht = request.form.get('tipo_ht', '1')
                condicion = request.form.get('condicion')
                estado_equipo = request.form.get('estado_equipo')
                deficiencia = request.form.get('deficiencia', '').strip()
                trabajo = request.form.get('trabajo', '').strip()
                observaciones = request.form.get('observaciones', '').strip()
                
                fecha_recepcion = request.form.get('fecha_recepcion')
                if not fecha_recepcion:
                    fecha_recepcion = date.today().strftime('%Y-%m-%d')
                fecha_entrega = request.form.get('fecha_entrega')
                if not fecha_entrega:
                    fecha_entrega = date.today().strftime('%Y-%m-%d')
                hora_ejecucion = request.form.get('hora_ejecucion', '12:00').strip()
                
                try:
                    tiempo_reparacion = float(request.form.get('tiempo_reparacion', '0').strip().replace(',', '.'))
                    if tiempo_reparacion < 0:
                        tiempo_reparacion = 0.0
                except:
                    tiempo_reparacion = 0.0
                
                repuesto_usado = request.form.get('repuesto_usado') == 'on'
                repuesto_nombre = request.form.get('repuesto_nombre', '').strip() if repuesto_usado else ''
                repuesto_cantidad_str = request.form.get('repuesto_cantidad', '0') if repuesto_usado else '0'
                repuesto_cantidad = int(repuesto_cantidad_str) if repuesto_cantidad_str.isdigit() else 0

                # Si el repuesto viene de la lista autocompletada, limpiar el sufijo de stock
                if repuesto_usado and " (Disponible:" in repuesto_nombre:
                    repuesto_nombre = repuesto_nombre.split(" (Disponible:")[0].strip()

                estado_final = request.form.get('estado_final', 'Operativo')

                detalle_final = f"Trabajo: {trabajo}."
                if deficiencia:
                    detalle_final = f"Deficiencia: {deficiencia}. " + detalle_final

                # 1. Registrar la intervención
                cur.execute("""
                    INSERT INTO historial_intervenciones (
                        equipo_id, fecha, tipo, detalle, condicion, estado_equipo,
                        deficiencia, trabajo, observaciones, fecha_entrega,
                        servicio_ht, tipo_ht, repuesto_usado, repuesto_nombre, repuesto_cantidad, realizado_por, hora_entrega, tiempo_reparacion
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (
                    id_equipo, fecha_recepcion, tipo, detalle_final, condicion, estado_equipo,
                    deficiencia, trabajo, observaciones, fecha_entrega,
                    eq['servicio'], tipo_ht, repuesto_usado, repuesto_nombre, repuesto_cantidad, usuario_valido['nombre_completo'], hora_ejecucion, tiempo_reparacion
                ))


                # 2. Descontar stock del repuesto si corresponde
                if repuesto_usado and repuesto_nombre:
                    cur.execute("""
                        UPDATE repuestos 
                        SET cantidad = GREATEST(0, cantidad - %s) 
                        WHERE nombre_repuesto = %s
                    """, (repuesto_cantidad, repuesto_nombre))

                # 3. Actualizar el estado del equipo
                cur.execute("UPDATE equipos SET estado = %s WHERE id = %s", (estado_final, id_equipo))
                
                conn.commit()
                cur.close()
                conn.close()

                # Generar archivos Excel y PDF de la Hoja de Trabajo
                form_data = {
                    'tipo': tipo, 'condicion': condicion, 'estado_equipo': estado_equipo,
                    'deficiencia': deficiencia, 'trabajo': trabajo, 'observaciones': observaciones,
                    'tipo_ht': tipo_ht, 'fecha_recepcion': fecha_recepcion, 'hora_ejecucion': hora_ejecucion
                }
                sello_firma_path = usuario_valido.get("sello_firma")
                
                xlsx_file, pdf_file = generar_excel_ht_web(dict(eq), form_data, usuario_valido['nombre_completo'], sello_firma_path)

                # Marcar los datos como sucios para sincronizar en tiempo real con la GUI del software
                global app_gui
                if app_gui:
                    app_gui.datos_sucios = True

                return render_template_string(HTML_EXITO, id_equipo=id_equipo, xlsx_file=xlsx_file, pdf_file=pdf_file)

        cur.close()
        conn.close()

        html_formulario = """
        <!DOCTYPE html><html lang="es"><head><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Registrar Mantenimiento</title>
        <style>
            *, *:before, *:after {
                box-sizing: border-box;
            }
            body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #F2F2F7; margin: 0; padding: 15px; color: #1C1C1E; }
            .tarjeta { background: #fff; border-radius: 16px; padding: 20px; box-shadow: 0 4px 24px rgba(0,0,0,0.06); max-width: 600px; margin: auto; }
            .cabecera { background: #007AFF; color: white; padding: 15px; border-radius: 12px; text-align: center; margin-bottom: 20px; }
            .campo { margin-bottom: 16px; }
            label { display: block; font-weight: bold; margin-bottom: 6px; font-size: 14px; color: #3A3A3C; }
            select, input[type="text"], input[type="password"], input[type="date"], input[type="time"], input[type="number"], textarea { 
                display: block;
                width: 100%; 
                padding: 12px; 
                border: 1px solid #C7C7CC; 
                border-radius: 8px; 
                font-size: 15px; 
                background: #FFF; 
                color: #1C1C1E; 
                margin: 0;
            }
            select:focus, input[type="text"]:focus, input[type="password"]:focus, input[type="date"]:focus, input[type="time"]:focus, input[type="number"]:focus, textarea:focus { outline: none; border-color: #007AFF; }
            .check-group { display: flex; align-items: center; gap: 10px; margin-top: 10px; }
            .check-group input { width: 20px; height: 20px; }
            .repuesto-caja { background: #F2F2F7; padding: 15px; border-radius: 10px; margin-top: 10px; display: none; }
            .btn-save { background: #34C759; color: white; font-weight: bold; border: none; width: 100%; padding: 15px; border-radius: 10px; font-size: 16px; margin-top: 10px; cursor: pointer; }
            .btn-save:active { background: #30B34F; }
            .btn-cancel { display: block; text-align: center; margin-top: 15px; color: #FF3B30; text-decoration: none; font-size: 15px; font-weight: bold; }
        </style>
        <script>
            function toggleRepuestos() {
                var checked = document.getElementById("repuesto_usado").checked;
                document.getElementById("caja-repuestos").style.display = checked ? "block" : "none";
            }

            function intentarSincronizarAhora() {
                var pendientes = JSON.parse(localStorage.getItem("mantenimientos_pendientes") || "[]");
                if (pendientes.length === 0) return;
                
                var status = document.getElementById("offline-sync-status");
                if (status) status.innerHTML = "🔄 Sincronizando reporte offline con el servidor... (" + pendientes.length + " restantes)";
                
                var item = pendientes[0];
                var passVal = item.web_pass;
                if (!passVal) {
                    passVal = prompt("Introduce tu contraseña (" + (item.web_user || "usuario") + ") para sincronizar:");
                    if (!passVal) {
                        if (status) status.innerHTML = "<span style='color:#FF9500;'>⚠️ Sincronización pausada: contraseña requerida.</span>";
                        return;
                    }
                }

                var params = new URLSearchParams();
                for (var key in item) {
                    if (key !== "web_pass") {
                        params.append(key, item[key]);
                    }
                }
                params.append("web_pass", passVal);
                
                fetch(item.url_sincronizacion, {
                    method: "POST",
                    body: params
                })
                .then(response => {
                    if (response.ok) {
                        pendientes.shift();
                        localStorage.setItem("mantenimientos_pendientes", JSON.stringify(pendientes));
                        if (pendientes.length > 0) {
                            intentarSincronizarAhora();
                        } else {
                            if (status) status.innerHTML = "<span style='color:#34C759;font-weight:bold;'>✅ ¡Reporte guardado y sincronizado con éxito!</span>";
                            response.text().then(html => {
                                setTimeout(() => {
                                    document.body.innerHTML = html;
                                    window.scrollTo(0, 0);
                                }, 1000);
                            });
                        }
                    } else {
                        response.text().then(txt => {
                            if (txt.includes("Usuario o contraseña incorrectos")) {
                                if (status) status.innerHTML = "<span style='color:#FF3B30;'>⚠️ Error: Credenciales incorrectas.</span>";
                            } else {
                                if (status) status.innerHTML = "<span style='color:#FF3B30;'>⚠️ El servidor rechazó la sincronización.</span>";
                            }
                        });
                    }
                })
                .catch(err => {
                    if (status) status.innerHTML = "⚠️ El servidor sigue fuera de línea. Reintentando al recuperar señal...";
                });
            }

            function guardarMantenimientoLocal(data) {
                // Por seguridad hospitalaria: NUNCA persistir la contraseña en localStorage
                delete data.web_pass;
                
                var pendientes = JSON.parse(localStorage.getItem("mantenimientos_pendientes") || "[]");
                data.id_equipo = "{{ eq['id'] }}";
                data.nombre_equipo = "{{ eq['nombre'] }}";
                data.fecha_registro_offline = new Date().toISOString();
                data.url_sincronizacion = window.location.href;
                pendientes.push(data);
                localStorage.setItem("mantenimientos_pendientes", JSON.stringify(pendientes));

                
                // Mostrar pantalla de éxito offline
                document.body.innerHTML = `
                <div class="tarjeta" style="max-width: 450px; text-align: center; padding: 30px; margin: 40px auto;">
                    <div style="font-size: 50px; color: #FF9500; margin-bottom: 15px;">⚠️</div>
                    <h2 style="margin-top:0; color: #E08200;">Registro Guardado Offline</h2>
                    <p style="color: #666; font-size: 14px; margin-bottom: 25px;">
                        No hay señal o conexión al servidor del hospital. El reporte de mantenimiento para <strong>${data.nombre_equipo}</strong> se ha guardado localmente en tu celular.
                    </p>
                    <div style="background: #FFF9E6; border: 1px solid #FFE0B2; padding: 12px; border-radius: 8px; font-size: 13px; color: #B78103; margin-bottom: 20px; text-align: left;">
                        💡 <strong>¿Qué hacer ahora?</strong><br>
                        En cuanto recuperes la red o Wi-Fi del hospital, mantén esta pestaña del navegador abierta. El celular enviará de forma automática el reporte de manera directa.
                    </div>
                    <button class="btn-save" onclick="intentarSincronizarAhora()" style="background: #007AFF;">🔄 Intentar Sincronizar Ahora</button>
                    <div id="offline-sync-status" style="margin-top: 15px; font-size: 13px; font-weight: bold; color: #666;"></div>
                </div>
                `;
            }

            window.addEventListener("load", function() {
                var pendientes = JSON.parse(localStorage.getItem("mantenimientos_pendientes") || "[]");
                if (pendientes.length > 0) {
                    intentarSincronizarAhora();
                }

                // Interceptar envío del formulario para soportar modo sin conexión
                var form = document.querySelector("form");
                if (form) {
                    form.addEventListener("submit", function(e) {
                        e.preventDefault();
                        
                        var btn = document.querySelector(".btn-save");
                        btn.disabled = true;
                        btn.innerText = "Procesando...";

                        var formData = new FormData(form);
                        var data = {};
                        formData.forEach((value, key) => { data[key] = value });
                        
                        // Asegurar checkboxes
                        if (!data.repuesto_usado) data.repuesto_usado = "";
                        else data.repuesto_usado = "on";

                        var params = new URLSearchParams(formData);

                        fetch(window.location.href, {
                            method: "POST",
                            body: params
                        })
                        .then(response => {
                            if (response.ok) {
                                response.text().then(html => {
                                    document.body.innerHTML = html;
                                    window.scrollTo(0, 0);
                                });
                            } else {
                                response.text().then(html => {
                                    document.body.innerHTML = html;
                                    window.scrollTo(0, 0);
                                });
                            }
                        })
                        .catch(err => {
                            // Falla de red/conexión, almacenar offline
                            guardarMantenimientoLocal(data);
                        });
                    });
                }
            });
        </script></head><body>
            <div class="tarjeta">
                <div class="cabecera">
                    <h3 style="margin:0;">Nueva Intervención</h3>
                    <span style="font-size:12px;">{{ eq['nombre'] }} (ID: {{ eq['id'] }})</span>
                </div>
                <form method="POST">
                    <!-- Sección de seguridad requerida -->
                    <div class="campo" style="background: #FFF9E6; padding: 15px; border-radius: 10px; border: 1px solid #FFE0B2; margin-bottom: 20px;">
                        <label style="color: #FF9500; font-size: 13px;">🔒 Validar Credenciales (Seguridad)</label>
                        <div style="margin-bottom: 10px;">
                            <input type="text" name="web_user" required placeholder="Usuario">
                        </div>
                        <div>
                            <input type="password" name="web_pass" required placeholder="Contraseña">
                        </div>
                        {% if error %}
                        <div style="color: #FF3B30; font-size: 13px; font-weight: bold; margin-top: 10px;">⚠️ {{ error }}</div>
                        {% endif %}
                    </div>

                    <div class="campo">
                        <label>Tipo de Mantenimiento</label>
                        <select name="tipo">
                            <option value="Preventivo">🔧 Preventivo</option>
                            <option value="Correctivo">🚨 Correctivo</option>
                        </select>
                    </div>
                    <div class="campo">
                        <label>Fecha de Recepción</label>
                        <input type="date" name="fecha_recepcion" value="{{ hoy_str }}" required>
                    </div>
                    <div class="campo">
                        <label>Fecha de Entrega</label>
                        <input type="date" name="fecha_entrega" value="{{ hoy_str }}" required>
                    </div>
                    <div class="campo">
                        <label>Hora de Entrega</label>
                        <input type="time" name="hora_ejecucion" value="{{ hora_str }}" required>
                    </div>
                    <div class="campo">
                        <label>Tiempo de Reparación (Horas)</label>
                        <input type="number" name="tiempo_reparacion" step="0.1" min="0" value="0" required placeholder="0">
                    </div>
                    <div class="campo">
                        <label>Condición de Entrega</label>
                        <select name="condicion">
                            <option value="Óptimo">Óptimo</option>
                            <option value="Aceptable">Aceptable</option>
                            <option value="Crítica">Crítica</option>
                            <option value="Inoperante">Inoperante</option>
                            <option value="F/Servicio">Fuera de Servicio</option>
                        </select>
                    </div>
                    <div class="campo">
                        <label>Estado Físico del Equipo</label>
                        <select name="estado_equipo">
                            <option value="Óptimo">Óptimo</option>
                            <option value="Bueno">Bueno</option>
                            <option value="Regular">Regular</option>
                            <option value="Malo">Malo</option>
                            <option value="Obsoleto">Obsoleto</option>
                        </select>
                    </div>
                    <div class="campo">
                        <label>Estado Operativo Final (Para el Catálogo)</label>
                        <select name="estado_final">
                            <option value="Operativo">Operativo</option>
                            <option value="En Espera de Repuesto">En Espera de Repuesto</option>
                            <option value="Fuera de Servicio">Fuera de Servicio</option>
                            <option value="Baja">Dado de Baja</option>
                        </select>
                    </div>
                    <div class="campo">
                        <label>Deficiencia Reportada</label>
                        <textarea name="deficiencia" rows="2" placeholder="Describa el problema reportado..."></textarea>
                    </div>
                    <div class="campo">
                        <label>Trabajo Realizado</label>
                        <textarea name="trabajo" rows="3" required placeholder="Detalle las pruebas o reparaciones hechas..."></textarea>
                    </div>
                    <div class="campo">
                        <label>Observaciones</label>
                        <textarea name="observaciones" rows="2" placeholder="Notas adicionales..."></textarea>
                    </div>
                    
                    <div class="campo check-group">
                        <input type="checkbox" id="repuesto_usado" name="repuesto_usado" onchange="toggleRepuestos()">
                        <label for="repuesto_usado" style="margin:0;">¿Se utilizó repuesto?</label>
                    </div>
                    
                    <div id="caja-repuestos" class="repuesto-caja">
                        <div class="campo">
                            <label>Nombre del Repuesto</label>
                            <input type="text" list="repuestos_list" name="repuesto_nombre" placeholder="Busca o escribe el repuesto...">
                            <datalist id="repuestos_list">
                                {% for r in repuestos %}
                                <option value="{{ r['nombre_repuesto'] }} (Disponible: {{ r['cantidad'] }})">
                                {% endfor %}
                            </datalist>
                        </div>
                        <div class="campo">
                            <label>Cantidad Utilizada</label>
                            <input type="number" name="repuesto_cantidad" min="1" value="1">
                        </div>
                    </div>
                    
                    <button type="submit" class="btn-save">Guardar Registro</button>
                    <a href="/equipo/{{ eq['id'] }}" class="btn-cancel">Cancelar</a>
                </form>
            </div>
        </body></html>
        """
        hoy_str = date.today().strftime('%Y-%m-%d')
        hora_str = datetime.now().strftime('%H:%M')
        return render_template_string(html_formulario, eq=eq, error=error_msg, repuestos=repuestos_list, hoy_str=hoy_str, hora_str=hora_str)
    except Exception as e:
        return f"Error en el servidor web: {e}"

# =========================================================================
# RUTAS DE AUTENTICACIÓN MÓVIL Y GESTIÓN DE SESIÓN
# =========================================================================

@app_web.route('/login', methods=['GET', 'POST'])
@app_web.route('/movil/login', methods=['GET', 'POST'])
def movil_login():
    """Pantalla de acceso con usuario y contraseña para el registro móvil en celulares."""
    error = None
    next_url = request.args.get('next') or url_for('movil_registro')
    
    if request.method == 'POST':
        u = request.form.get('usuario', '').strip()
        p = request.form.get('password', '').strip()
        
        if not u or not p:
            error = "Por favor ingrese su usuario y contraseña."
        else:
            user_info = auth_login(u, p)
            if user_info:
                session['usuario'] = user_info['nombre_usuario']
                session['nombre_completo'] = user_info.get('nombre_completo', user_info['nombre_usuario'])
                session['rol'] = user_info.get('rol', 'tecnico')
                # Cargar y parsear permisos del usuario
                raw_perms = user_info.get('permisos') or {}
                if isinstance(raw_perms, str):
                    try: raw_perms = json.loads(raw_perms)
                    except: raw_perms = {}
                session['permisos'] = raw_perms
                session.permanent = True
                return redirect(next_url)
            else:
                error = "Usuario o contraseña incorrectos. Verifique sus credenciales."
                
    elif 'usuario' in session:
        return redirect(next_url)
        
    return render_template_string(HTML_MOVIL_LOGIN, error=error)

@app_web.route('/logout')
@app_web.route('/movil/logout')
def movil_logout():
    """Cierra la sesión del usuario actual."""
    session.clear()
    return redirect(url_for('movil_login'))

# =========================================================================
# RUTAS DE REGISTRO MÓVIL DESDE CELULARES
# =========================================================================

@app_web.route('/movil')
@app_web.route('/movil/registro')
@login_requerido
def movil_registro():
    """Formulario interactivo y táctil para registro de equipamiento médico desde celular."""
    try:
        sedes_data = obtener_jerarquia_sedes_db()
        usuario_dict = {
            "nombre_usuario": session.get('usuario', ''),
            "nombre_completo": session.get('nombre_completo', session.get('usuario', 'Técnico')),
            "rol": session.get('rol', 'tecnico'),
            "permisos": session.get('permisos', {})
        }
        return render_template_string(
            HTML_MOVIL_REGISTRO,
            sedes=sedes_data,
            usuario=usuario_dict
        )
    except Exception as e:
        return f"Error cargando formulario móvil: {e}", 500

# =========================================================================
# API REST JSON PARA LA APLICACIÓN MÓVIL Y HEALTHCHECK / PING
# =========================================================================

@app_web.route('/ping')
@app_web.route('/api/health')
def api_ping_health():
    """Keep-alive ultra liviano para UptimeRobot / Cron.
    Mantiene despiertos simultáneamente Render y la base de datos Supabase."""
    db_ok = False
    error_msg = None
    try:
        conn = obtener_conexion()
        if conn:
            with conn.cursor() as cur:
                cur.execute("SELECT 1;")
            conn.close()
            db_ok = True
    except Exception as e:
        error_msg = str(e)

    status_code = 200 if db_ok else 503
    return jsonify({
        "status": "ok" if db_ok else "degraded",
        "service": "CMMS-GAMLP Web Server (Render)",
        "database": "Supabase PostgreSQL (Activo)" if db_ok else f"DB Error: {error_msg}",
        "timestamp": datetime.now().isoformat()
    }), status_code

# =========================================================================
# ENDPOINTS JSON DE CONSULTA HISTÓRICA (SOLO LECTURA • GESTIONES ANTERIORES)
# =========================================================================

@app_web.route('/api/historico/sedes')
def api_historico_sedes():
    """Retorna la jerarquía de sedes de la base histórica."""
    try:
        data = obtener_jerarquia_sedes_db(perfil='historica')
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/historico/equipos')
def api_historico_equipos():
    """Retorna equipos médicos de la base histórica (solo lectura)."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        equipos = obtener_equipos_db(centro_nombre=centro or None, limite=5000, perfil='historica', red_nombre=red or None)
        return jsonify(equipos)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/historico/muebles')
def api_historico_muebles():
    """Retorna activos de mueblería y computación de la base histórica (solo lectura)."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        muebles = obtener_muebles_db(centro_nombre=centro or None, limite=5000, perfil='historica', red_nombre=red or None)
        return jsonify(muebles)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/historico/areas')
def api_historico_areas():
    """Retorna áreas registradas en la base histórica."""
    centro = request.args.get('centro', '').strip()
    try:
        areas = obtener_areas_db(centro_nombre=centro or None, perfil='historica')
        return jsonify(areas)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/historico/catalogo')
def api_historico_catalogo():
    """Retorna catálogo de modelos de la base histórica."""
    try:
        cats = obtener_catalogo_equipos_db(perfil='historica')
        return jsonify(cats)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/historico/estadisticas')
def api_historico_estadisticas():
    """Retorna estadísticas e indicadores calculados sobre la base histórica."""
    centro = request.args.get('centro', '').strip()
    try:
        stats = obtener_estadisticas_censo_db(centro_nombre=centro or None, perfil='historica')
        return jsonify(stats)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/sedes')
def api_sedes():
    """Retorna la jerarquía de redes y centros de salud en formato JSON."""
    try:
        data = obtener_jerarquia_sedes_db()
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/areas')
def api_areas():
    """Retorna las áreas registradas, opcionalmente filtradas por centro de salud y red."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        areas = obtener_areas_db(centro_nombre=centro or None, red_nombre=red or None)
        return jsonify(areas)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/catalogo')
def api_catalogo():
    """Retorna la lista de modelos de equipos preconfigurados en el catálogo."""
    try:
        catalogo = obtener_catalogo_equipos_db()
        return jsonify(catalogo)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/siguiente_af')
def api_siguiente_af():
    """Calcula y retorna el siguiente código correlativo de Activos Fijos (AF)."""
    red = request.args.get('red', '').strip()
    centro = request.args.get('centro', '').strip()
    if not red or not centro:
        return jsonify({"ok": False, "error": "Parámetros 'red' y 'centro' requeridos"}), 400
    try:
        codigo = generar_siguiente_codigo_af(red, centro)
        return jsonify({"ok": True, "codigo_af": codigo})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/guardar_equipo', methods=['POST'])
@login_requerido
def api_guardar_equipo():
    """Guarda un nuevo equipo médico registrado desde el celular con resolución de colisiones."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True)
        if not payload:
            return jsonify({"ok": False, "error": "Datos inválidos o cuerpo de solicitud vacío"}), 400

        nom = str(payload.get("nombre") or "").strip()
        if not nom:
            return jsonify({"ok": False, "error": "El nombre del equipo es obligatorio."}), 400

        # Guardar en base de datos PostgreSQL con resolución de colisión automática
        exito, id_final, msg = guardar_equipo_db(payload)

        # Si la ventana de escritorio de Tkinter está abierta, refrescarla en tiempo real
        if app_gui:
            try:
                app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except Exception as ge:
                print(f"[DEBUG] No se pudo refrescar GUI de escritorio: {ge}")

        return jsonify({
            "ok": exito,
            "id": id_final,
            "nombre": nom,
            "mensaje": msg
        })
    except Exception as e:
        print(f"[ERROR] Error en api_guardar_equipo: {e}")
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/equipos')
def api_equipos():
    """Retorna la lista de equipos médicos registrados."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        equipos = obtener_equipos_db(centro_nombre=centro or None, limite=3000, red_nombre=red or None)
        return jsonify(equipos)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/guardar_catalogo', methods=['POST'])
@login_requerido
def api_guardar_catalogo():
    """Guarda un nuevo modelo en el catálogo central."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True) or {}
        if not payload.get('nombre'):
            return jsonify({"ok": False, "error": "El nombre del modelo es obligatorio"}), 400
        ok, c_id = guardar_catalogo_db(payload)
        if ok and app_gui:
            try: app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except: pass
        return jsonify({"ok": ok, "id": c_id})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/muebles')
def api_muebles():
    """Retorna los activos de mueblería y computación."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        muebles = obtener_muebles_db(centro_nombre=centro or None, limite=3000, red_nombre=red or None)
        return jsonify(muebles)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/guardar_mueble', methods=['POST'])
@login_requerido
def api_guardar_mueble():
    """Guarda o actualiza un activo de mueblería o cómputo."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True) or {}
        if not payload.get('tipo_activo') and not payload.get('descripcion'):
            return jsonify({"ok": False, "error": "Tipo de activo o descripción requerida"}), 400
        
        if not payload.get('descripcion'):
            payload['descripcion'] = payload.get('tipo_activo') or 'Activo Mueblería / TI'

        # Obtener nombre del técnico de la sesión de forma segura
        nom_tecnico = ''
        u_movil = session.get('usuario_movil')
        if isinstance(u_movil, dict):
            nom_tecnico = u_movil.get('nombre_completo') or u_movil.get('nombre_usuario') or ''
        elif session.get('nombre_completo'):
            nom_tecnico = str(session.get('nombre_completo'))
        elif session.get('usuario'):
            nom_tecnico = str(session.get('usuario'))
        nom_tecnico = nom_tecnico.strip()

        if not payload.get('tecnico_inventareador') and nom_tecnico:
            payload['tecnico_inventareador'] = nom_tecnico

        ok, res = guardar_mueble_db(payload)
        if not ok:
            return jsonify({"ok": False, "error": str(res)}), 400

        if app_gui:
            try: app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except: pass
        return jsonify({"ok": True, "id": res})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/guardar_area', methods=['POST'])
@login_requerido
def api_guardar_area():
    """Guarda una nueva área por centro de salud."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True) or {}
        if not payload.get('nombre'):
            return jsonify({"ok": False, "error": "El nombre del área es obligatorio"}), 400
        ok, msg = guardar_area_db(payload)
        if ok and app_gui:
            try: app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except: pass
        return jsonify({"ok": ok, "mensaje": msg})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/repuestos')
def api_repuestos():
    """Retorna la lista de repuestos en stock."""
    centro = request.args.get('centro', '').strip()
    try:
        repuestos = obtener_repuestos_db(centro)
        return jsonify(repuestos)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/guardar_repuesto', methods=['POST'])
@login_requerido
def api_guardar_repuesto():
    """Guarda o actualiza un repuesto en la base de datos."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True) or {}
        if not payload.get('nombre_repuesto'):
            return jsonify({"ok": False, "error": "El nombre del repuesto es obligatorio"}), 400
        ok, msg = guardar_repuesto_db(payload)
        if ok and app_gui:
            try: app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except: pass
        return jsonify({"ok": ok, "mensaje": msg})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/usuarios')
@login_requerido
def api_usuarios():
    """Retorna la lista de usuarios con sus permisos (Solo administradores / jefes)."""
    rol = str(session.get('rol', '')).strip().lower()
    u = str(session.get('usuario', '')).strip().lower()
    if rol not in ['admin', 'administrador', 'jefe'] and u != 'godhead':
        return jsonify({"error": "Acceso restringido a administradores."}), 403
    try:
        users = obtener_usuarios_db()
        return jsonify(users)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/guardar_usuario_permisos', methods=['POST'])
@login_requerido
def api_guardar_usuario_permisos():
    """Guarda o actualiza un usuario, su rol y su matriz de permisos (Solo administradores / jefes)."""
    rol = str(session.get('rol', '')).strip().lower()
    u = str(session.get('usuario', '')).strip().lower()
    if rol not in ['admin', 'administrador', 'jefe'] and u != 'godhead':
        return jsonify({"ok": False, "error": "Acceso restringido a administradores."}), 403
    try:
        payload = request.get_json(force=True, silent=True) or {}
        ok, msg = guardar_usuario_permisos_db(payload)
        return jsonify({"ok": ok, "mensaje": msg})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/intervenciones')
@login_requerido
def api_intervenciones():
    """Retorna historial de intervenciones y mantenimientos técnicos."""
    centro = request.args.get('centro', '').strip()
    eq_id = request.args.get('equipo_id', '').strip()
    try:
        data = obtener_intervenciones_db(centro_nombre=centro, equipo_id=eq_id)
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/guardar_intervencion', methods=['POST'])
@login_requerido
def api_guardar_intervencion():
    """Registra una intervención técnica y actualiza el estado del equipo."""
    global app_gui
    try:
        payload = request.get_json(force=True, silent=True) or {}
        if not payload.get('equipo_id'):
            return jsonify({"ok": False, "error": "El código del equipo es obligatorio."}), 400
        
        # Asignar técnico actual si no viene en el payload
        if not payload.get('realizado_por'):
            payload['realizado_por'] = session.get('nombre_completo', session.get('usuario', 'Técnico'))

        ok, res = guardar_intervencion_db(payload)
        if ok and app_gui:
            try: app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
            except: pass
        return jsonify({"ok": ok, "id": res if ok else None, "mensaje": res if not ok else "Intervención guardada"})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app_web.route('/api/estadisticas')
@login_requerido
def api_estadisticas():
    """Retorna indicadores y métricas de censo de equipos."""
    centro = request.args.get('centro', '').strip()
    red = request.args.get('red', '').strip()
    try:
        stats = obtener_estadisticas_censo_db(centro or None, red_nombre=red or None)
        return jsonify(stats)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/eliminar_registro', methods=['POST'])
@login_requerido
def api_eliminar_registro():
    """Elimina de forma segura un registro de cualquier módulo (con verificación de permisos y papelera)."""
    global app_gui
    rol = str(session.get('rol', '')).strip().lower()
    u_ci = str(session.get('usuario', '')).strip().lower()
    es_admin = (rol in ['admin', 'administrador', 'jefe']) or (u_ci == 'godhead')

    payload = request.get_json(force=True, silent=True) or {}
    tabla = str(payload.get('tabla') or '').strip().lower()
    id_reg = payload.get('id')

    if not tabla or not id_reg:
        return jsonify({"ok": False, "error": "Parámetros 'tabla' e 'id' son requeridos."}), 400

    # Mapeo de tabla a módulo para control de permisos
    modulo_map = {
        "equipos": "Inventario",
        "catalogo": "Catalogo",
        "muebleria": "Muebleria",
        "areas": "Areas",
        "repuestos": "Repuestos",
        "historial_intervenciones": "Historial",
        "usuarios": "Usuarios"
    }
    mod_key = modulo_map.get(tabla)
    if not es_admin:
        perms = session.get('permisos') or {}
        puede_eliminar = perms.get(mod_key, {}).get('eliminar', False) if mod_key else False
        if not puede_eliminar:
            return jsonify({"ok": False, "error": f"No tiene permisos para eliminar registros en {mod_key}."}), 403

    usr = session.get('nombre_completo', session.get('usuario', 'web_user'))
    ok, msg = eliminar_registro_db(tabla, id_reg, usuario=usr)

    if ok and app_gui:
        try:
            app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
        except Exception:
            pass

    return jsonify({"ok": ok, "mensaje": msg})

@app_web.route('/api/papelera')
@login_requerido
def api_papelera():
    """Retorna los elementos en papelera. Exclusivo para administradores y godhead."""
    rol = str(session.get('rol', '')).strip().lower()
    u = str(session.get('usuario', '')).strip().lower()
    if rol not in ['admin', 'administrador', 'jefe'] and u != 'godhead':
        return jsonify({"error": "Acceso denegado. Exclusivo para modo Administrador y Godhead."}), 403

    tabla = request.args.get('tabla', '').strip().lower()
    q = request.args.get('q', '').strip()
    try:
        items = obtener_papelera_db(filtro_tabla=tabla or None, busqueda=q or None, limite=300)
        return jsonify(items)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app_web.route('/api/papelera/recuperar', methods=['POST'])
@login_requerido
def api_papelera_recuperar():
    """Restaura un elemento de la papelera a su tabla origen. Exclusivo para admin y godhead."""
    global app_gui
    rol = str(session.get('rol', '')).strip().lower()
    u = str(session.get('usuario', '')).strip().lower()
    if rol not in ['admin', 'administrador', 'jefe'] and u != 'godhead':
        return jsonify({"ok": False, "error": "Acceso denegado. Exclusivo para modo Administrador y Godhead."}), 403

    payload = request.get_json(force=True, silent=True) or {}
    p_id = payload.get('id')
    if not p_id:
        return jsonify({"ok": False, "error": "Parámetro 'id' requerido."}), 400

    try:
        p_id = int(p_id)
    except Exception:
        return jsonify({"ok": False, "error": "ID de papelera inválido."}), 400

    usr = session.get('nombre_completo', session.get('usuario', 'admin'))
    ok, msg = recuperar_registro_papelera_db(p_id, usuario=usr)
    if ok and app_gui:
        try:
            app_gui.after(300, app_gui.cargar_datos_en_segundo_plano)
        except Exception:
            pass

    return jsonify({"ok": ok, "mensaje": msg})

@app_web.route('/api/papelera/eliminar_definitivo', methods=['POST'])
@login_requerido
def api_papelera_eliminar_definitivo():
    """Elimina definitivamente un elemento de la papelera. Exclusivo para admin y godhead."""
    rol = str(session.get('rol', '')).strip().lower()
    u = str(session.get('usuario', '')).strip().lower()
    if rol not in ['admin', 'administrador', 'jefe'] and u != 'godhead':
        return jsonify({"ok": False, "error": "Acceso denegado. Exclusivo para modo Administrador y Godhead."}), 403

    payload = request.get_json(force=True, silent=True) or {}
    p_id = payload.get('id')
    vaciar = bool(payload.get('vaciar_todo', False))

    if not p_id and not vaciar:
        return jsonify({"ok": False, "error": "Parámetro 'id' o 'vaciar_todo' requerido."}), 400

    try:
        p_id_int = int(p_id) if p_id else None
    except Exception:
        p_id_int = None

    ok, msg = purgar_registro_papelera_db(papelera_id=p_id_int, vaciar_todo=vaciar)
    return jsonify({"ok": ok, "mensaje": msg})

def iniciar_servidor_web():
    import logging
    logging.getLogger('werkzeug').setLevel(logging.ERROR)
    app_web.run(host='0.0.0.0', port=5000, debug=False, use_reloader=False, threaded=True)


def arrancar_hilo_web(gui_instance=None):
    """Función para ser llamada desde el archivo principal que levanta el servidor en segundo plano."""
    global app_gui
    if gui_instance:
        app_gui = gui_instance
    hilo_web = threading.Thread(target=iniciar_servidor_web, daemon=True)
    hilo_web.start()
    return hilo_web

if __name__ == '__main__':
    import socket as _sock

    def obtener_ip_local():
        try:
            s = _sock.socket(_sock.AF_INET, _sock.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            s.close()
            return ip
        except:
            return "127.0.0.1"

    def puerto_en_uso(puerto):
        with _sock.socket(_sock.AF_INET, _sock.SOCK_STREAM) as s:
            return s.connect_ex(('127.0.0.1', puerto)) == 0

    ip = obtener_ip_local()
    print("="*60)
    print("     SGEM GAMLP - SERVIDOR WEB INDEPENDIENTE (v1.1)")
    print(f" Servidor activo en: http://{ip}:5000")
    print(" Mantén esta ventana abierta para que los códigos QR funcionen")
    print(" incluso cuando el programa principal esté cerrado.")
    print("="*60)

    if puerto_en_uso(5000):
        print("\n[AVISO] El puerto 5000 ya está en uso (probablemente el software principal ya arrancó el servidor).")
        print(" El servidor web ya está funcionando correctamente.")
        input("\nPresiona Enter para cerrar esta ventana...")
    else:
        iniciar_servidor_web()
